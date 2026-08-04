# -*- coding: utf-8 -*-
"""
ServerCRON — agendador de scripts Python lidos de registro_automacoes.xlsx.

Layout esperado (raiz do projeto = pasta do server.py):

  ServerCRON/
    server.py
    install_deps.bat
    registro_automacoes.xlsx   <-- planilha neste nivel (obrigatorio)
    automacoes/
      area/.../script.py

Planilha (colunas):
  nome_automacao | cron_schedule | is_active

nome_automacao = nome do ficheiro .py a correr
  (ex.: ``backup_diario`` -> ``python backup_diario.py``).

Arranque:
  python server.py
  (corre install_deps.bat -> pip install -> reinicia e sobe o servidor)
"""
from __future__ import annotations

import calendar
import csv
import hmac
import logging
import json
import os
import re
import secrets
import signal
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Optional
from zoneinfo import ZoneInfo

# ---------------------------------------------------------------------------
# Boot: optional .env + pip sync
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent


def _load_dotenv(base: Path) -> None:
    path = base / ".env"
    if not path.is_file():
        return
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return
    for line in raw.splitlines():
        s = line.strip()
        if not s or s.startswith("#") or "=" not in s:
            continue
        key, _, val = s.partition("=")
        key = key.strip()
        if not key:
            continue
        val = val.strip()
        if len(val) >= 2 and val[0] == val[-1] and val[0] in "\"'":
            val = val[1:-1]
        if key not in os.environ:
            os.environ[key] = val


_load_dotenv(BASE_DIR)


def _run_install_deps_bat() -> None:
    """Corre install_deps.bat (python -m pip install -r requirements.txt)."""
    bat = BASE_DIR / "install_deps.bat"
    req = BASE_DIR / "requirements.txt"
    if not bat.is_file():
        print(f"[BOOT] Ficheiro em falta: {bat}", file=sys.stderr, flush=True)
        raise SystemExit(1)
    if not req.is_file():
        print(f"[BOOT] Ficheiro em falta: {req}", file=sys.stderr, flush=True)
        raise SystemExit(1)

    env = os.environ.copy()
    env["SERVERCRON_PYTHON"] = sys.executable
    print("[BOOT] A correr install_deps.bat …", flush=True)
    try:
        subprocess.check_call(["cmd.exe", "/c", str(bat)], cwd=str(BASE_DIR), env=env)
    except subprocess.CalledProcessError as exc:
        print("[BOOT] Falha ao instalar dependencias via .bat", file=sys.stderr, flush=True)
        raise SystemExit(1) from exc


def _reexec_after_deps() -> None:
    """Reinicia este processo para carregar as libraries ja instaladas."""
    env = os.environ.copy()
    env["SERVERCRON_SKIP_REQUIREMENTS_PIP"] = "1"
    script = str(Path(__file__).resolve())
    args = [sys.executable, script, *sys.argv[1:]]
    print("[BOOT] A recarregar libraries e a iniciar o servidor…", flush=True)
    os.execve(sys.executable, args, env)


if __name__ == "__main__" and str(os.environ.get("SERVERCRON_SKIP_REQUIREMENTS_PIP", "")).lower() not in (
    "1",
    "true",
    "yes",
):
    _run_install_deps_bat()
    _reexec_after_deps()

from croniter import croniter  # noqa: E402
from flask import Flask, jsonify, request, send_from_directory  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from waitress import serve  # noqa: E402

# ---------------------------------------------------------------------------
# Config — caminhos sempre via Path.home() (portavel em qualquer Windows)
# ---------------------------------------------------------------------------

def _env_path(name: str, default: Path) -> Path:
    """Le path de env; expande ~ e resolve. Sem hardcode de C:\\Users\\..."""
    raw = (os.environ.get(name) or "").strip()
    if not raw:
        return default.expanduser().resolve()
    return Path(raw).expanduser().resolve()


def _resolve_user_path(raw: str) -> Path:
    return Path(str(raw or "").strip()).expanduser().resolve()


HOME = Path.home()
DATA_ROOT = _env_path("SERVERCRON_DATA_ROOT", HOME / "Documents" / "ServerCRON")
CONFIG_PATH = DATA_ROOT / "config.json"
# .py ficam em automacoes/; a planilha fica SEMPRE na raiz do projeto
# (mesmo nivel de server.py e install_deps.bat) — nao dentro de automacoes/
DEFAULT_AUTOMAOES_DIR = BASE_DIR / "automacoes"
DEFAULT_REGISTRO_XLSX = (BASE_DIR / "registro_automacoes.xlsx").resolve()
# Historico SQLite / logs / config ficam no home do utilizador
SQLITE_PATH = DATA_ROOT / "server_cron.sqlite"
LOGS_DIR = DATA_ROOT / "logs"
HISTORY_CSV_PATH = LOGS_DIR / "historico_execucoes.csv"

# Mutaveis: pasta das automacoes (painel / config.json / env)
AUTOMAOES_DIR: Path = DEFAULT_AUTOMAOES_DIR
# Planilha: raiz do projeto por defeito. Env so se quiseres override explicito.
_reg_env = (os.environ.get("SERVERCRON_REGISTRO_XLSX") or "").strip()
REGISTRO_XLSX: Path = (
    Path(_reg_env).expanduser().resolve() if _reg_env else DEFAULT_REGISTRO_XLSX
)
_automacoes_from_env = bool((os.environ.get("SERVERCRON_AUTOMAOES_DIR") or "").strip())

_registry_lock = threading.Lock()
_registry_cache: list[dict[str, Any]] = []
_registry_mtime: float = 0.0
_registry_ts: float = 0.0

_users_lock = threading.Lock()
_users_cache: list[dict[str, Any]] = []
_users_mtime: float = 0.0
_users_ts: float = 0.0

# Sessoes emitidas apos login por email (token::{email, expires_at}).
# Em memoria: reiniciar o server derruba todas as sessoes (aceitavel, o login e rapido).
_sessions_lock = threading.Lock()
_sessions: dict[str, dict[str, Any]] = {}
SESSION_TTL_SEC = max(60, int((os.environ.get("SERVERCRON_SESSION_TTL_SEC") or str(10 * 3600)).strip() or str(10 * 3600)))

# Rate limit de tentativas de login por IP (mitiga forca bruta de emails).
_login_attempts_lock = threading.Lock()
_login_attempts: dict[str, list[float]] = {}
LOGIN_MAX_ATTEMPTS = 8
LOGIN_WINDOW_SEC = 300

# Codigo de acesso enviado por email (OTP): prova posse da caixa de entrada,
# nao so do endereco. Uso unico, expira rapido, poucas tentativas de verificacao.
_otp_lock = threading.Lock()
_otp_codes: dict[str, dict[str, Any]] = {}  # email -> {"code", "expires_at", "attempts"}
OTP_TTL_SEC = max(30, int((os.environ.get("SERVERCRON_OTP_TTL_SEC") or "120").strip() or "120"))
OTP_CODE_BYTES = 18  # secrets.token_urlsafe(18) -> ~24 caracteres, alta entropia
OTP_MAX_VERIFY_ATTEMPTS = 5

SMTP_HOST = (os.environ.get("SERVERCRON_SMTP_HOST") or "smtp.gmail.com").strip()
SMTP_PORT = int((os.environ.get("SERVERCRON_SMTP_PORT") or "465").strip() or "465")
SMTP_USER = (os.environ.get("SERVERCRON_SMTP_USER") or "").strip()
SMTP_PASSWORD = (os.environ.get("SERVERCRON_SMTP_PASSWORD") or "").strip()


def _smtp_configured() -> bool:
    return bool(SMTP_USER and SMTP_PASSWORD)


def _generate_otp() -> str:
    return secrets.token_urlsafe(OTP_CODE_BYTES)


def _send_otp_email(to_email: str, code: str) -> None:
    """Envia o codigo por SMTP (Gmail com App Password por padrao). Lanca excecao se falhar
    (o chamador decide o que fazer — a rota trata e devolve 502 ao utilizador)."""
    import smtplib
    from email.message import EmailMessage

    minutos = max(1, OTP_TTL_SEC // 60)
    msg = EmailMessage()
    msg["Subject"] = "Codigo de acesso - ServerCRON"
    msg["From"] = SMTP_USER
    msg["To"] = to_email
    msg.set_content(
        "Seu codigo de acesso ao ServerCRON:\n\n"
        f"{code}\n\n"
        f"Expira em {minutos} minuto(s) e so pode ser usado uma vez.\n"
        "Se voce nao pediu este codigo, pode ignorar este email."
    )
    with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=15) as smtp:
        smtp.login(SMTP_USER, SMTP_PASSWORD)
        smtp.send_message(msg)


def _read_config_file() -> dict[str, Any]:
    if not CONFIG_PATH.is_file():
        return {}
    try:
        data = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _write_config_file(data: dict[str, Any]) -> None:
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def apply_automacoes_dir(path: Path, *, persist: bool = False) -> Path:
    """Define a pasta dos .py (a planilha permanece na raiz do projeto)."""
    global AUTOMAOES_DIR, _registry_cache, _registry_mtime, _registry_ts
    resolved = path.expanduser().resolve()
    AUTOMAOES_DIR = resolved
    with _registry_lock:
        _registry_cache = []
        _registry_mtime = 0.0
        _registry_ts = 0.0
    if persist and not _automacoes_from_env:
        cfg = _read_config_file()
        cfg["automacoes_dir"] = str(AUTOMAOES_DIR)
        _write_config_file(cfg)
    log.info("[CONFIG] automacoes_dir=%s | planilha=%s", AUTOMAOES_DIR, REGISTRO_XLSX)
    return AUTOMAOES_DIR


def load_automacoes_dir_from_sources() -> Path:
    """Prioridade: env SERVERCRON_AUTOMAOES_DIR > config.json > pasta do projeto/automacoes."""
    env_raw = (os.environ.get("SERVERCRON_AUTOMAOES_DIR") or "").strip()
    if env_raw:
        return apply_automacoes_dir(_resolve_user_path(env_raw), persist=False)
    cfg = _read_config_file()
    cfg_raw = str(cfg.get("automacoes_dir") or "").strip()
    if cfg_raw:
        return apply_automacoes_dir(_resolve_user_path(cfg_raw), persist=False)
    return apply_automacoes_dir(DEFAULT_AUTOMAOES_DIR, persist=False)


HOST = (os.environ.get("SERVERCRON_HOST") or "127.0.0.1").strip() or "127.0.0.1"
PORT = int((os.environ.get("SERVERCRON_PORT") or "5001").strip() or "5001")
OPEN_BROWSER = str(os.environ.get("SERVERCRON_OPEN_BROWSER", "1")).lower() in ("1", "true", "yes")
RELOAD_MINUTES = max(1, int((os.environ.get("SERVERCRON_RELOAD_MINUTES") or "5").strip() or "5"))
# Fuso fixo das automacoes: horario de Brasilia (Sao Paulo). Override so via SERVERCRON_TZ.
TZ_NAME = (os.environ.get("SERVERCRON_TZ") or "America/Sao_Paulo").strip() or "America/Sao_Paulo"
MAX_CONCURRENT = max(1, int((os.environ.get("SERVERCRON_MAX_CONCURRENT") or "5").strip() or "5"))
# Quantos slots em atraso por automacao enfileirar de cada vez (do mais antigo)
CATCHUP_PER_SCRIPT = max(1, int((os.environ.get("SERVERCRON_CATCHUP_PER_SCRIPT") or "5").strip() or "5"))
# Proximos horarios do dia (ainda por vir) a manter na fila por automacao
UPCOMING_PER_SCRIPT = max(1, int((os.environ.get("SERVERCRON_UPCOMING_PER_SCRIPT") or "12").strip() or "12"))
# Clima Sao Paulo via Open-Meteo (open source, sem API key)
WEATHER_CACHE_SEC = max(60, int((os.environ.get("SERVERCRON_WEATHER_CACHE_SEC") or "900").strip() or "900"))
WEATHER_LAT = float((os.environ.get("SERVERCRON_WEATHER_LAT") or "-23.5505").strip() or "-23.5505")
WEATHER_LON = float((os.environ.get("SERVERCRON_WEATHER_LON") or "-46.6333").strip() or "-46.6333")
# Token opcional para proteger /api/*. Vazio = painel aberto (ok em 127.0.0.1 uso pessoal).
# Obrigatorio se HOST nao for loopback (ver checagem no arranque, main()).
API_TOKEN = (os.environ.get("SERVERCRON_API_TOKEN") or "").strip()
# Kill switch TEMPORARIO para demo/apresentacao: desliga toda a exigencia de login
# (token estatico + OTP por email), inclusive a guarda de exposicao na rede.
# Nao apaga nenhuma logica de auth -- so a ignora enquanto a env var estiver ligada.
# Remover SERVERCRON_DISABLE_AUTH do .env (ou por "0"/vazio) restaura a protecao normal.
AUTH_DISABLED = (os.environ.get("SERVERCRON_DISABLE_AUTH") or "").strip().lower() in (
    "1",
    "true",
    "yes",
)
# Mata a automacao se ela passar deste tempo (evita 1 script preso a ocupar 1 slot para sempre). 0 = sem limite.
JOB_TIMEOUT_SEC = max(0, int((os.environ.get("SERVERCRON_JOB_TIMEOUT_SEC") or "3600").strip() or "3600"))
# Intervalo minimo entre disparos manuais da MESMA automacao pelo painel (anti double-click / spam).
MANUAL_RUN_COOLDOWN_SEC = max(0, float((os.environ.get("SERVERCRON_MANUAL_RUN_COOLDOWN_SEC") or "2").strip() or "2"))

REQUIRED_COLUMNS = ("nome_automacao", "cron_schedule", "is_active")
CRON_RE = re.compile(
    r"^(\S+)\s+(\S+)\s+(\S+)\s+(\S+)\s+(\S+)(?:\s+(\S+))?$"
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("servercron")

# Ultima leitura forçada da planilha (epoch). Usado para UI /api/status.
_last_spreadsheet_reload_ts: float = 0.0
_spreadsheet_reload_lock = threading.Lock()

def mark_spreadsheet_reload() -> None:
    """Regista o momento da ultima leitura da planilha + rebuild de crons."""
    global _last_spreadsheet_reload_ts
    with _spreadsheet_reload_lock:
        _last_spreadsheet_reload_ts = time.time()


def spreadsheet_reload_meta() -> dict[str, Any]:
    """Ultima e proxima atualizacao automatica dos crons da planilha."""
    with _spreadsheet_reload_lock:
        last_ts = float(_last_spreadsheet_reload_ts or 0.0)
    interval = RELOAD_MINUTES * 60
    if last_ts <= 0:
        return {
            "reload_minutes": RELOAD_MINUTES,
            "last_reload_at": None,
            "next_reload_at": None,
            "seconds_until_reload": None,
        }
    next_ts = last_ts + interval
    now = time.time()
    tz = ZoneInfo(TZ_NAME)
    return {
        "reload_minutes": RELOAD_MINUTES,
        "last_reload_at": datetime.fromtimestamp(last_ts, tz=tz).isoformat(timespec="seconds"),
        "next_reload_at": datetime.fromtimestamp(next_ts, tz=tz).isoformat(timespec="seconds"),
        "seconds_until_reload": max(0, int(next_ts - now)),
    }


# ---------------------------------------------------------------------------
# Time helpers — sempre America/Sao_Paulo (horario de Brasilia)
# ---------------------------------------------------------------------------

_TZ_SP: ZoneInfo | None = None


def _tz() -> ZoneInfo:
    """Timezone do cron: America/Sao_Paulo. Nunca usa o fuso do Windows do PC."""
    global _TZ_SP
    if _TZ_SP is not None:
        return _TZ_SP
    try:
        _TZ_SP = ZoneInfo(TZ_NAME)
    except Exception:
        # Windows sem tzdata: tenta de novo com nome canonico; ultimo recurso UTC + aviso
        try:
            _TZ_SP = ZoneInfo("America/Sao_Paulo")
        except Exception:
            log.error(
                "[TZ] Nao foi possivel carregar America/Sao_Paulo. "
                "Instale o pacote tzdata (pip install tzdata). A usar UTC temporariamente."
            )
            _TZ_SP = ZoneInfo("UTC")
    return _TZ_SP


def _to_sp(dt: datetime) -> datetime:
    """Converte qualquer datetime para America/Sao_Paulo."""
    if dt.tzinfo is None:
        return dt.replace(tzinfo=_tz())
    return dt.astimezone(_tz())


def now_local() -> datetime:
    """Agora no horario de Brasilia (Sao Paulo)."""
    return datetime.now(_tz())


def now_iso() -> str:
    return now_local().isoformat(timespec="seconds")


def _croniter_sp(cron_expr: str, base: Optional[datetime] = None) -> croniter:
    """croniter ancorado no horario de Sao Paulo."""
    base_sp = _to_sp(base or now_local())
    return croniter(cron_expr, base_sp)


# ---------------------------------------------------------------------------
# Excel registry
# ---------------------------------------------------------------------------


def _norm_col(name: Any) -> str:
    return str(name or "").strip().lower().replace(" ", "_")


def _parse_bool(val: Any) -> bool:
    if isinstance(val, bool):
        return val
    s = str(val or "").strip().upper()
    return s in ("TRUE", "1", "YES", "Y", "SIM", "S", "ATIVO", "ACTIVE")


def _safe_name(val: Any) -> str:
    return str(val or "").strip()


def _script_stem(val: Any) -> str:
    """nome_automacao = nome do ficheiro .py a executar (com ou sem extensao).

    Ex.: planilha ``backup_diario`` ou ``backup_diario.py`` -> procura ``backup_diario.py``.
    """
    nome = _safe_name(val)
    if nome.lower().endswith(".py"):
        nome = nome[:-3].strip()
    return nome


def _is_valid_cron(expr: str) -> bool:
    expr = (expr or "").strip()
    if not expr or not CRON_RE.match(expr):
        return False
    try:
        _croniter_sp(expr, now_local())
        return True
    except (ValueError, KeyError, TypeError):
        return False


def _maybe_migrate_registro_xlsx() -> None:
    """Move/copia a planilha para a raiz do projeto se ainda estiver noutro sitio."""
    if REGISTRO_XLSX.is_file():
        return
    import shutil

    candidates = [
        AUTOMAOES_DIR / "registro_automacoes.xlsx",
        BASE_DIR / "automacoes" / "registro_automacoes.xlsx",
        HOME / "Documents" / "ServerCRON" / "automacoes" / "registro_automacoes.xlsx",
        HOME / "Documents" / "ServerCRON" / "registro_automacoes.xlsx",
    ]
    for src in candidates:
        if not src.is_file():
            continue
        if src.resolve() == REGISTRO_XLSX.resolve():
            return
        try:
            REGISTRO_XLSX.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, REGISTRO_XLSX)
            log.info(
                "[BOOT] Planilha movida para o nivel do server.py: %s (origem %s)",
                REGISTRO_XLSX,
                src,
            )
            if src.parent.resolve() in (
                AUTOMAOES_DIR.resolve(),
                (BASE_DIR / "automacoes").resolve(),
            ):
                try:
                    src.unlink(missing_ok=True)
                except OSError:
                    pass
            return
        except OSError:
            log.exception("[BOOT] Falha ao copiar planilha de %s", src)


def ensure_registro_at_project_root() -> Path:
    """Garante que a planilha usada e a da raiz do projeto (nivel do server.py)."""
    global REGISTRO_XLSX
    expected = (BASE_DIR / "registro_automacoes.xlsx").resolve()
    # sem override por env: forca o caminho canonico
    if not (os.environ.get("SERVERCRON_REGISTRO_XLSX") or "").strip():
        REGISTRO_XLSX = expected
    _maybe_migrate_registro_xlsx()
    if not REGISTRO_XLSX.is_file():
        # cria vazia/exemplo na raiz — nunca dentro de automacoes/
        wb = Workbook()
        ws = wb.active
        ws.title = "AUTOMACOES"
        ws.append(list(REQUIRED_COLUMNS))
        ws.append(["exemplo_hello", "*/5 * * * *", "FALSE"])
        REGISTRO_XLSX.parent.mkdir(parents=True, exist_ok=True)
        wb.save(REGISTRO_XLSX)
        log.info("[BOOT] Criada planilha em %s (nivel do server.py)", REGISTRO_XLSX)
    log.info(
        "[BOOT] Planilha esperada no mesmo nivel de server.py: %s (existe=%s)",
        REGISTRO_XLSX,
        REGISTRO_XLSX.is_file(),
    )
    return REGISTRO_XLSX


def _maybe_migrate_automacoes_from_documents() -> None:
    """Se a pasta do projeto ainda nao tem .py, copia de Documents/ServerCRON/automacoes."""
    has_py = AUTOMAOES_DIR.is_dir() and any(AUTOMAOES_DIR.rglob("*.py"))
    if has_py:
        return
    legacy = HOME / "Documents" / "ServerCRON" / "automacoes"
    if not legacy.is_dir():
        return
    import shutil

    log.info("[BOOT] A migrar scripts de %s -> %s", legacy, AUTOMAOES_DIR)
    AUTOMAOES_DIR.mkdir(parents=True, exist_ok=True)
    for item in legacy.iterdir():
        if item.name.lower() == "registro_automacoes.xlsx":
            continue
        dest = AUTOMAOES_DIR / item.name
        if dest.exists():
            continue
        try:
            if item.is_dir():
                shutil.copytree(item, dest)
            elif item.is_file():
                shutil.copy2(item, dest)
        except OSError:
            log.exception("[BOOT] Falha ao copiar %s", item)


def ensure_automacoes_layout() -> None:
    """Cria pasta automacoes + area exemplo; planilha fica na raiz (nivel do server.py)."""
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    AUTOMAOES_DIR.mkdir(parents=True, exist_ok=True)
    _maybe_migrate_automacoes_from_documents()
    ensure_registro_at_project_root()
    area_exemplo = AUTOMAOES_DIR / "exemplo"
    area_exemplo.mkdir(parents=True, exist_ok=True)
    sample = area_exemplo / "exemplo_hello.py"
    if not sample.is_file():
        sample.write_text(
            'print("ServerCRON exemplo_hello OK")\n',
            encoding="utf-8",
        )
        log.info("[BOOT] Criado %s", sample)


def list_area_dirs() -> list[str]:
    """Nomes/caminhos relativos das pastas (sob automacoes/) que contem .py."""
    if not AUTOMAOES_DIR.is_dir():
        return []
    areas: set[str] = set()
    for py in AUTOMAOES_DIR.rglob("*.py"):
        if not py.is_file():
            continue
        try:
            rel = py.parent.resolve().relative_to(AUTOMAOES_DIR.resolve())
        except ValueError:
            continue
        areas.add("." if str(rel) == "." else str(rel).replace("\\", "/"))
    return sorted(areas, key=str.lower)


def find_script(nome: str) -> tuple[Optional[Path], Optional[str]]:
    """Procura ``{nome}.py`` em toda a arvore de automacoes/ (recursivo).

    A planilha ``registro_automacoes.xlsx`` fica na raiz do projeto (junto de
    ``server.py``). Os scripts ficam em ``automacoes/`` (areas e subpastas).
    Retorna (path, area_relativa).
    """
    stem = _script_stem(nome)
    if not stem or not AUTOMAOES_DIR.is_dir():
        return None, None
    key = stem.lower()
    root = AUTOMAOES_DIR.resolve()
    matches: list[tuple[Path, str]] = []
    for py in AUTOMAOES_DIR.rglob("*.py"):
        if not py.is_file():
            continue
        if py.stem.lower() != key:
            continue
        try:
            parent_rel = py.parent.resolve().relative_to(root)
        except ValueError:
            continue
        area = "." if str(parent_rel) == "." else str(parent_rel).replace("\\", "/")
        matches.append((py.resolve(), area))
    if not matches:
        return None, None
    # Preferir caminho mais curto (mais perto da raiz); depois ordem alfabetica
    matches.sort(key=lambda t: (t[0].as_posix().count("/"), t[1].lower(), t[0].as_posix().lower()))
    return matches[0]


def find_script_path(nome: str) -> Optional[Path]:
    path, _ = find_script(nome)
    return path


def scan_python_files() -> list[dict[str, str]]:
    """Lista recursiva de todos os .py sob automacoes/."""
    if not AUTOMAOES_DIR.is_dir():
        return []
    root = AUTOMAOES_DIR.resolve()
    out: list[dict[str, str]] = []
    for py in AUTOMAOES_DIR.rglob("*.py"):
        if not py.is_file():
            continue
        try:
            rel_parent = py.parent.resolve().relative_to(root)
        except ValueError:
            continue
        area = "." if str(rel_parent) == "." else str(rel_parent).replace("\\", "/")
        out.append(
            {
                "nome": py.stem,
                "area": area,
                "path": str(py.resolve()),
            }
        )
    out.sort(key=lambda x: (x["area"].lower(), x["nome"].lower()))
    return out


def inventory_summary(scripts: Optional[list[dict[str, Any]]] = None) -> dict[str, Any]:
    """Resumo para o painel: disco vs planilha, ativos/desativados, localizados."""
    if scripts is None:
        scripts = read_registry()
    py_files = scan_python_files()
    active = [s for s in scripts if s.get("is_active")]
    inactive = [s for s in scripts if not s.get("is_active")]
    located = [s for s in scripts if s.get("available_locally")]
    missing = [s for s in scripts if not s.get("available_locally")]
    active_located = [s for s in active if s.get("available_locally")]
    return {
        "py_on_disk": len(py_files),
        "areas_count": len(list_area_dirs()),
        "sheet_total": len(scripts),
        "sheet_active": len(active),
        "sheet_inactive": len(inactive),
        "sheet_located": len(located),
        "sheet_missing": len(missing),
        "sheet_active_located": len(active_located),
        "schedulable": len(schedulable()),
        "registro_exists": REGISTRO_XLSX.is_file(),
        "automacoes_dir": str(AUTOMAOES_DIR),
        "py_files": py_files,
    }


def read_registry(force: bool = False) -> list[dict[str, Any]]:
    global _registry_cache, _registry_mtime, _registry_ts
    with _registry_lock:
        mtime = REGISTRO_XLSX.stat().st_mtime if REGISTRO_XLSX.is_file() else 0.0
        if (
            not force
            and _registry_cache
            and mtime == _registry_mtime
            and (time.time() - _registry_ts) < 30
        ):
            return [dict(r) for r in _registry_cache]

    if not REGISTRO_XLSX.is_file():
        log.warning("[REGISTRY] Planilha ausente: %s", REGISTRO_XLSX)
        with _registry_lock:
            _registry_cache = []
            _registry_mtime = 0.0
            _registry_ts = time.time()
        return []

    try:
        wb = load_workbook(REGISTRO_XLSX, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            wb.close()
            return []
        headers = [_norm_col(h) for h in header_row]
        col_map = {h: i for i, h in enumerate(headers) if h}
        missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
        if missing:
            log.error(
                "[REGISTRY] Colunas em falta %s. Encontrado: %s",
                missing,
                headers,
            )
            wb.close()
            return []

        records: list[dict[str, Any]] = []
        for row in rows:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            def cell(col: str) -> Any:
                idx = col_map[col]
                return row[idx] if idx < len(row) else None

            nome = _script_stem(cell("nome_automacao"))
            if not nome:
                continue
            cron_raw = _safe_name(cell("cron_schedule"))
            active = _parse_bool(cell("is_active"))
            script, area = find_script(nome)
            next_run = None
            valid = _is_valid_cron(cron_raw)
            if active and valid:
                try:
                    nxt = _croniter_sp(cron_raw, now_local()).get_next(datetime)
                    next_run = _to_sp(nxt).isoformat(timespec="seconds")
                except Exception:
                    next_run = None
            records.append(
                {
                    "nome_automacao": nome,
                    "area": area,
                    "cron_schedule": cron_raw,
                    "is_active": active,
                    "is_valid_cron": valid,
                    "available_locally": script is not None,
                    "path": str(script) if script else None,
                    "next_run": next_run,
                }
            )
        wb.close()
    except Exception:
        log.exception("[REGISTRY] Erro ao ler %s", REGISTRO_XLSX)
        with _registry_lock:
            return [dict(r) for r in _registry_cache]

    with _registry_lock:
        _registry_cache = records
        _registry_mtime = REGISTRO_XLSX.stat().st_mtime if REGISTRO_XLSX.is_file() else 0.0
        _registry_ts = time.time()
    log.info("[REGISTRY] %d automacao(oes) em %s", len(records), REGISTRO_XLSX)
    return [dict(r) for r in records]


USERS_REQUIRED_COLUMNS = ("email",)


def read_users(force: bool = False) -> list[dict[str, Any]]:
    """Le a aba USERS da planilha (email, nome, is_active). Aba ausente = lista vazia."""
    global _users_cache, _users_mtime, _users_ts
    with _users_lock:
        mtime = REGISTRO_XLSX.stat().st_mtime if REGISTRO_XLSX.is_file() else 0.0
        if (
            not force
            and _users_cache
            and mtime == _users_mtime
            and (time.time() - _users_ts) < 30
        ):
            return [dict(r) for r in _users_cache]

    if not REGISTRO_XLSX.is_file():
        with _users_lock:
            _users_cache = []
            _users_mtime = 0.0
            _users_ts = time.time()
        return []

    try:
        wb = load_workbook(REGISTRO_XLSX, read_only=True, data_only=True)
        sheet_name = next((s for s in wb.sheetnames if s.strip().upper() == "USERS"), None)
        if not sheet_name:
            wb.close()
            with _users_lock:
                _users_cache = []
                _users_mtime = mtime
                _users_ts = time.time()
            return []
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            wb.close()
            return []
        headers = [_norm_col(h) for h in header_row]
        col_map = {h: i for i, h in enumerate(headers) if h}
        missing = [c for c in USERS_REQUIRED_COLUMNS if c not in col_map]
        if missing:
            log.error("[USERS] Colunas em falta %s na aba USERS. Encontrado: %s", missing, headers)
            wb.close()
            return []

        records: list[dict[str, Any]] = []
        for row in rows:
            if not row or all(c is None or str(c).strip() == "" for c in row):
                continue

            def cell(col: str) -> Any:
                idx = col_map.get(col)
                return row[idx] if idx is not None and idx < len(row) else None

            email = _safe_name(cell("email")).strip().lower()
            if not email or "@" not in email:
                continue
            records.append(
                {
                    "email": email,
                    "nome": _safe_name(cell("nome")) or None,
                    "is_active": _parse_bool(cell("is_active")) if "is_active" in col_map else True,
                    "is_admin": _parse_bool(cell("is_admin")) if "is_admin" in col_map else False,
                }
            )
        wb.close()
    except Exception:
        log.exception("[USERS] Erro ao ler aba USERS de %s", REGISTRO_XLSX)
        with _users_lock:
            return [dict(r) for r in _users_cache]

    with _users_lock:
        _users_cache = records
        _users_mtime = REGISTRO_XLSX.stat().st_mtime if REGISTRO_XLSX.is_file() else 0.0
        _users_ts = time.time()
    return [dict(r) for r in records]


def _email_norm(raw: Any) -> str:
    return _safe_name(raw).strip().lower()


def _email_allowed(raw_email: Any) -> bool:
    email = _email_norm(raw_email)
    if not email or "@" not in email:
        return False
    return any(u["email"] == email and u["is_active"] for u in read_users())


def _email_is_admin(raw_email: Any) -> bool:
    email = _email_norm(raw_email)
    if not email:
        return False
    return any(u["email"] == email and u["is_active"] and u["is_admin"] for u in read_users())


def write_users(records: list[dict[str, Any]]) -> None:
    """Reescreve a aba USERS por completo (usado pelo painel de gestao de usuarios)."""
    global _users_cache, _users_mtime, _users_ts
    with _users_lock:
        if REGISTRO_XLSX.is_file():
            wb = load_workbook(REGISTRO_XLSX)
        else:
            wb = Workbook()
            default = wb.active
            default.title = "AUTOMACOES"
            default.append(list(REQUIRED_COLUMNS))
        if "USERS" in wb.sheetnames:
            del wb["USERS"]
        ws = wb.create_sheet("USERS")
        ws.append(["email", "nome", "is_active", "is_admin"])
        for r in records:
            ws.append(
                [
                    r["email"],
                    r.get("nome") or "",
                    "TRUE" if r.get("is_active", True) else "FALSE",
                    "TRUE" if r.get("is_admin", False) else "FALSE",
                ]
            )
        REGISTRO_XLSX.parent.mkdir(parents=True, exist_ok=True)
        wb.save(REGISTRO_XLSX)
        wb.close()
        _users_cache = []
        _users_mtime = 0.0
        _users_ts = 0.0


def _cleanup_sessions_locked() -> None:
    now = time.time()
    expired = [tok for tok, s in _sessions.items() if s["expires_at"] <= now]
    for tok in expired:
        del _sessions[tok]


def issue_session(email: str) -> dict[str, Any]:
    """Emite uma sessao nova para o email e revoga qualquer sessao anterior do MESMO email
    (uma sessao ativa por pessoa — logar de novo noutro PC derruba a sessao antiga)."""
    token = secrets.token_urlsafe(32)
    expires_at = time.time() + SESSION_TTL_SEC
    with _sessions_lock:
        _cleanup_sessions_locked()
        for tok, s in list(_sessions.items()):
            if s["email"] == email:
                del _sessions[tok]
        _sessions[token] = {"email": email, "expires_at": expires_at}
    return {"token": token, "expires_at": expires_at}


def session_email(token: str) -> Optional[str]:
    if not token:
        return None
    with _sessions_lock:
        session = _sessions.get(token)
        if not session:
            return None
        if session["expires_at"] <= time.time():
            del _sessions[token]
            return None
        return session["email"]


def revoke_session(token: str) -> None:
    with _sessions_lock:
        _sessions.pop(token, None)


def _login_rate_limited(ip: str) -> bool:
    now = time.time()
    with _login_attempts_lock:
        attempts = [t for t in _login_attempts.get(ip, []) if now - t < LOGIN_WINDOW_SEC]
        _login_attempts[ip] = attempts
        return len(attempts) >= LOGIN_MAX_ATTEMPTS


def _register_login_attempt(ip: str) -> None:
    with _login_attempts_lock:
        _login_attempts.setdefault(ip, []).append(time.time())


# Rate limit generico por (bucket, ip) — protege /api/run, /api/kill, /api/catchup
# contra flood (spam de requisicoes derrubando o server / esgotando disco+CPU),
# independente do cooldown por automacao que ja existe em /api/run.
_rate_limit_lock = threading.Lock()
_rate_limit_hits: dict[str, list[float]] = {}


def _rate_limited(bucket: str, ip: str, max_hits: int, window_sec: float) -> bool:
    key = f"{bucket}:{ip}"
    now = time.time()
    with _rate_limit_lock:
        hits = [t for t in _rate_limit_hits.get(key, []) if now - t < window_sec]
        hits.append(now)
        _rate_limit_hits[key] = hits
        return len(hits) > max_hits


def schedulable(force: bool = False) -> list[dict[str, Any]]:
    out = []
    for r in read_registry(force=force):
        if not r["is_active"]:
            continue
        if not r["is_valid_cron"]:
            log.warning("[IGNORADO] %s: cron invalido (%s)", r["nome_automacao"], r["cron_schedule"])
            continue
        if not r["available_locally"]:
            log.warning("[IGNORADO] %s: ficheiro .py nao encontrado", r["nome_automacao"])
            continue
        out.append(r)
    return out


# ---------------------------------------------------------------------------
# History (SQLite)
# ---------------------------------------------------------------------------

_db_lock = threading.Lock()


def _db() -> sqlite3.Connection:
    conn = sqlite3.connect(str(SQLITE_PATH), check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    DATA_ROOT.mkdir(parents=True, exist_ok=True)
    with _db_lock:
        conn = _db()
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS execution_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome_automacao TEXT NOT NULL,
                area TEXT,
                trigger_reason TEXT NOT NULL,
                status TEXT NOT NULL,
                start_time TEXT NOT NULL,
                end_time TEXT,
                duration_sec REAL,
                exit_code INTEGER,
                output TEXT
            )
            """
        )
        cols = {
            str(r[1]).lower()
            for r in conn.execute("PRAGMA table_info(execution_runs)").fetchall()
        }
        if "area" not in cols:
            conn.execute("ALTER TABLE execution_runs ADD COLUMN area TEXT")
        conn.commit()
        conn.close()
    ensure_history_csv()


def insert_run(entry: dict[str, Any]) -> int:
    with _db_lock:
        conn = _db()
        cur = conn.execute(
            """
            INSERT INTO execution_runs
            (nome_automacao, area, trigger_reason, status, start_time, end_time, duration_sec, exit_code, output)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                entry.get("nome_automacao"),
                entry.get("area") or "",
                entry.get("trigger_reason", "manual"),
                entry.get("status", "running"),
                entry.get("start_time", now_iso()),
                entry.get("end_time"),
                entry.get("duration_sec"),
                entry.get("exit_code"),
                entry.get("output", ""),
            ),
        )
        conn.commit()
        rid = int(cur.lastrowid)
        # SQLite so para o painel (ultimas 500). O CSV e o historico longo (Power BI / Excel).
        conn.execute(
            """
            DELETE FROM execution_runs WHERE id NOT IN (
                SELECT id FROM execution_runs ORDER BY id DESC LIMIT 500
            )
            """
        )
        conn.commit()
        conn.close()
        return rid


def update_run(rid: int, **fields: Any) -> None:
    if not fields:
        return
    cols = ", ".join(f"{k}=?" for k in fields)
    vals = list(fields.values()) + [rid]
    with _db_lock:
        conn = _db()
        conn.execute(f"UPDATE execution_runs SET {cols} WHERE id=?", vals)
        conn.commit()
        row = conn.execute("SELECT * FROM execution_runs WHERE id=?", (rid,)).fetchone()
        conn.close()
    # Grava no CSV quando a execucao termina (nao em "running")
    if row and str(fields.get("status") or row["status"] or "") != "running":
        append_history_csv(dict(row))


# CSV permanente sob Path.home()/Documents/ServerCRON/logs — pronto para Power BI / Excel.
# Colunas de calendário + flags 0/1 facilitam filtros: 7 dias, 1 mês, 6 meses, etc.
_CSV_FIELDS = (
    "id",
    "nome_automacao",
    "area",
    "trigger_reason",
    "status",
    "exit_code",
    "duration_sec",
    "start_time",
    "end_time",
    "data",
    "hora",
    "ano",
    "mes",
    "ano_mes",
    "dia",
    "dia_semana",
    "dia_semana_num",
    "semana_ano",
    "flag_sucesso",
    "flag_erro",
    "flag_sem_dados",
    "output",
)
_CSV_WEEKDAYS = (
    "segunda",
    "terca",
    "quarta",
    "quinta",
    "sexta",
    "sabado",
    "domingo",
)
_csv_lock = threading.Lock()


def _parse_run_datetime(raw: Any) -> Optional[datetime]:
    text = str(raw or "").strip()
    if not text:
        return None
    try:
        dt = datetime.fromisoformat(text)
    except ValueError:
        return None
    return _to_sp(dt)


def csv_row_from_entry(entry: dict[str, Any]) -> dict[str, Any]:
    """Linha CSV enriquecida para dashboards (Power BI / Excel)."""
    out = entry.get("output") or ""
    if isinstance(out, str):
        out = out.replace("\r\n", "\\n").replace("\n", "\\n").replace("\r", "\\n")
        if len(out) > 4000:
            out = out[:3980] + "... [truncado]"
    status = str(entry.get("status") or "").strip().lower()
    start_raw = entry.get("start_time") or ""
    dt = _parse_run_datetime(start_raw)
    if dt is None:
        data = str(start_raw)[:10] if start_raw else ""
        hora = ""
        ano = mes = dia = semana = ""
        ano_mes = ""
        dia_semana = ""
        dia_semana_num = ""
        if len(data) == 10 and data[4] == "-" and data[7] == "-":
            ano, mes, dia = data[0:4], data[5:7].lstrip("0") or data[5:7], data[8:10].lstrip("0") or data[8:10]
            ano_mes = data[0:7]
    else:
        data = dt.strftime("%Y-%m-%d")
        hora = dt.strftime("%H:%M:%S")
        ano = dt.year
        mes = dt.month
        dia = dt.day
        ano_mes = dt.strftime("%Y-%m")
        dia_semana_num = dt.weekday() + 1  # 1=segunda ... 7=domingo
        dia_semana = _CSV_WEEKDAYS[dt.weekday()]
        semana = int(dt.strftime("%V"))
    return {
        "id": entry.get("id", ""),
        "nome_automacao": entry.get("nome_automacao", ""),
        "area": entry.get("area") or "",
        "trigger_reason": entry.get("trigger_reason", ""),
        "status": entry.get("status", ""),
        "exit_code": entry.get("exit_code", ""),
        "duration_sec": entry.get("duration_sec", ""),
        "start_time": start_raw,
        "end_time": entry.get("end_time", ""),
        "data": data,
        "hora": hora,
        "ano": ano,
        "mes": mes,
        "ano_mes": ano_mes,
        "dia": dia,
        "dia_semana": dia_semana,
        "dia_semana_num": dia_semana_num,
        "semana_ano": semana,
        "flag_sucesso": 1 if status == "success" else 0,
        "flag_erro": 1 if status == "error" else 0,
        "flag_sem_dados": 1 if status == "no_data" else 0,
        "output": out,
    }


def _csv_header_matches(path: Path) -> bool:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, None)
    except (OSError, StopIteration, csv.Error):
        return False
    if not header:
        return False
    return list(header) == list(_CSV_FIELDS)


def ensure_history_csv() -> Path:
    """Garante CSV sob Path.home() com cabecalho BI; migra formato antigo se preciso."""
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    with _csv_lock:
        if HISTORY_CSV_PATH.is_file() and _csv_header_matches(HISTORY_CSV_PATH):
            return HISTORY_CSV_PATH
        if HISTORY_CSV_PATH.is_file():
            stamp = now_local().strftime("%Y%m%d_%H%M%S")
            backup = LOGS_DIR / f"historico_execucoes_legado_{stamp}.csv"
            try:
                HISTORY_CSV_PATH.replace(backup)
                log.info("[CSV] Formato antigo arquivado em %s", backup)
                _migrate_legacy_history_csv(backup, HISTORY_CSV_PATH)
                return HISTORY_CSV_PATH
            except OSError:
                log.exception("[CSV] Falha ao migrar historico antigo")
        try:
            with HISTORY_CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=_CSV_FIELDS, extrasaction="ignore")
                writer.writeheader()
            log.info(
                "[CSV] Historico pronto para Power BI/Excel: %s",
                HISTORY_CSV_PATH,
            )
        except OSError:
            log.exception("[CSV] Falha ao criar %s", HISTORY_CSV_PATH)
    return HISTORY_CSV_PATH


def _migrate_legacy_history_csv(src: Path, dest: Path) -> None:
    """Reescreve CSV antigo no formato BI (colunas data/mes/flags)."""
    rows_out: list[dict[str, Any]] = []
    try:
        with src.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for raw in reader:
                rows_out.append(csv_row_from_entry(raw))
    except OSError:
        log.exception("[CSV] Nao foi possivel ler legado %s", src)
        with dest.open("w", encoding="utf-8-sig", newline="") as f:
            csv.DictWriter(f, fieldnames=_CSV_FIELDS, extrasaction="ignore").writeheader()
        return
    with dest.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=_CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for row in rows_out:
            writer.writerow(row)
    log.info("[CSV] Migradas %d linhas para formato BI em %s", len(rows_out), dest)


def append_history_csv(entry: dict[str, Any]) -> None:
    """Acrescenta uma execucao terminada ao CSV (nunca trunca; ideal para BI)."""
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    row = csv_row_from_entry(entry)
    with _csv_lock:
        try:
            if not HISTORY_CSV_PATH.is_file() or not _csv_header_matches(HISTORY_CSV_PATH):
                # fora do lock interno de ensure: ja estamos no lock
                if HISTORY_CSV_PATH.is_file() and not _csv_header_matches(HISTORY_CSV_PATH):
                    stamp = now_local().strftime("%Y%m%d_%H%M%S")
                    backup = LOGS_DIR / f"historico_execucoes_legado_{stamp}.csv"
                    try:
                        HISTORY_CSV_PATH.replace(backup)
                        _migrate_legacy_history_csv(backup, HISTORY_CSV_PATH)
                    except OSError:
                        log.exception("[CSV] falha na migracao lazy")
                if not HISTORY_CSV_PATH.is_file():
                    with HISTORY_CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
                        csv.DictWriter(f, fieldnames=_CSV_FIELDS, extrasaction="ignore").writeheader()
            with HISTORY_CSV_PATH.open("a", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=_CSV_FIELDS, extrasaction="ignore")
                writer.writerow(row)
        except OSError:
            log.exception("[CSV] falha ao gravar %s", HISTORY_CSV_PATH)


def list_history(limit: int = 100) -> list[dict[str, Any]]:
    with _db_lock:
        conn = _db()
        rows = conn.execute(
            "SELECT * FROM execution_runs ORDER BY id DESC LIMIT ?",
            (max(1, min(limit, 500)),),
        ).fetchall()
        conn.close()
    return [dict(r) for r in rows]


def list_history_today() -> list[dict[str, Any]]:
    """Execucoes de hoje (local), a partir do CSV — nao do SQLite.

    O SQLite guarda so as ultimas 500 corridas NO TOTAL (todas as automacoes, ver
    insert_run). Num dia com bastante catchup (ex.: um script "*/1 * * * *" parado
    de manha e recuperando a tarde) isso poda os horarios mais cedo do dia da tabela
    antes mesmo do fim do dia. Como detect_pending_today() usa esta lista pra saber
    se um horario ja foi coberto por alguma execucao, perder essas linhas fazia o
    painel mostrar "atrasada Xh" para sempre num horario que ja tinha rodado horas
    atras — so que o registro dele tinha sido descartado do SQLite nesse meio tempo.
    O CSV (historico_execucoes.csv) nunca e podado, entao e a fonte certa aqui.
    """
    today = now_local().strftime("%Y-%m-%d")
    if not HISTORY_CSV_PATH.is_file():
        return []
    out: list[dict[str, Any]] = []
    try:
        with HISTORY_CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for raw in reader:
                if str(raw.get("data") or "") == today:
                    out.append(raw)
    except OSError:
        log.exception("[PENDING] falha a ler CSV %s", HISTORY_CSV_PATH)
    return out


def execution_stats(*, days: int = 7) -> dict[str, Any]:
    """Agrega execucoes por dia e status a partir do CSV (fonte longa).

    O SQLite so guarda as ultimas 500 corridas e distorce o grafico do painel;
    o historico_execucoes.csv e a fonte correta para tendencia.
    """
    days = max(1, min(int(days), 30))
    tz_now = now_local()
    start_day = (tz_now - timedelta(days=days - 1)).date()
    end_day = tz_now.date()
    since = start_day.isoformat()
    until = end_day.isoformat()

    by_day: dict[str, dict[str, int]] = {}
    d = start_day
    while d <= end_day:
        key = d.isoformat()
        by_day[key] = {"success": 0, "error": 0, "no_data": 0, "total": 0}
        d += timedelta(days=1)

    totals = {"success": 0, "error": 0, "no_data": 0, "total": 0}
    per_script: dict[str, dict[str, Any]] = {}

    def _bump(day: str, st: str, nome: str) -> None:
        if day not in by_day or st not in ("success", "error", "no_data"):
            return
        by_day[day][st] += 1
        by_day[day]["total"] += 1
        totals[st] += 1
        totals["total"] += 1
        if not nome:
            return
        bucket = per_script.setdefault(
            nome,
            {"nome_automacao": nome, "success": 0, "error": 0, "no_data": 0, "total": 0},
        )
        bucket[st] += 1
        bucket["total"] += 1

    csv_used = False
    if HISTORY_CSV_PATH.is_file():
        try:
            with HISTORY_CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                for raw in reader:
                    day = str(raw.get("data") or "")[:10]
                    if not day or day < since or day > until:
                        continue
                    st = str(raw.get("status") or "").strip().lower()
                    nome = _script_stem(raw.get("nome_automacao"))
                    # catchup em massa distorce o grafico do dia; conta so scheduled/manual no chart
                    reason = str(raw.get("trigger_reason") or "").strip().lower()
                    if reason == "catchup":
                        continue
                    _bump(day, st, nome)
            csv_used = True
        except OSError:
            log.exception("[STATS] falha a ler CSV %s", HISTORY_CSV_PATH)

    if not csv_used or totals["total"] == 0:
        # fallback SQLite (painel ainda vazio / CSV em falta)
        with _db_lock:
            conn = _db()
            rows = conn.execute(
                """
                SELECT substr(start_time, 1, 10) AS day, status, nome_automacao, COUNT(*) AS n
                FROM execution_runs
                WHERE substr(start_time, 1, 10) >= ?
                  AND status IN ('success', 'error', 'no_data')
                  AND COALESCE(trigger_reason, '') != 'catchup'
                GROUP BY day, status, nome_automacao
                ORDER BY day ASC
                """,
                (since,),
            ).fetchall()
            conn.close()
        for r in rows:
            day = str(r["day"] or "")
            st = str(r["status"] or "")
            nome = _script_stem(r["nome_automacao"])
            n = int(r["n"] or 0)
            for _ in range(n):
                _bump(day, st, nome)

    today_key = end_day.isoformat()
    today = dict(by_day.get(today_key) or {"success": 0, "error": 0, "no_data": 0, "total": 0})
    top_scripts = sorted(per_script.values(), key=lambda x: x["total"], reverse=True)[:8]

    series = [
        {
            "day": day,
            "success": vals["success"],
            "error": vals["error"],
            "no_data": vals["no_data"],
            "total": vals["total"],
        }
        for day, vals in by_day.items()
    ]

    rate = error_rate = no_data_rate = None
    if totals["total"]:
        rate = round(100.0 * totals["success"] / totals["total"], 1)
        error_rate = round(100.0 * totals["error"] / totals["total"], 1)
        no_data_rate = round(100.0 * totals["no_data"] / totals["total"], 1)

    for item in top_scripts:
        t = int(item["total"] or 0)
        if t:
            item["success_pct"] = round(100.0 * item["success"] / t, 1)
            item["error_pct"] = round(100.0 * item["error"] / t, 1)
            item["no_data_pct"] = round(100.0 * item["no_data"] / t, 1)
        else:
            item["success_pct"] = None
            item["error_pct"] = None
            item["no_data_pct"] = None

    return {
        "days": days,
        "since": since,
        "until": until,
        "totals": totals,
        "today": today,
        "series": series,
        "top_scripts": top_scripts,
        "success_rate": rate,
        "error_rate": error_rate,
        "no_data_rate": no_data_rate,
        "rates": {
            "success": rate,
            "error": error_rate,
            "no_data": no_data_rate,
        },
        "source": "csv" if csv_used else "sqlite",
        "excludes_catchup": True,
    }


# ---------------------------------------------------------------------------
# Relatorio (aba "Relatorio" do painel) — totais desde o inicio, comparativos
# fixos (semana/mes) e drill-down por ano/mes/dia. Sempre a partir do CSV
# (fonte longa); exclui catchup, igual ao /api/stats.
# ---------------------------------------------------------------------------

_MESES_PT = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")


def _to_int_or_none(v: Any) -> Optional[int]:
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def _read_history_rows_no_catchup() -> list[dict[str, Any]]:
    if not HISTORY_CSV_PATH.is_file():
        return []
    out: list[dict[str, Any]] = []
    try:
        with HISTORY_CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for raw in reader:
                if str(raw.get("trigger_reason") or "").strip().lower() == "catchup":
                    continue
                out.append(raw)
    except OSError:
        log.exception("[REPORT] falha a ler CSV %s", HISTORY_CSV_PATH)
    return out


def _rows_kpi(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(rows)
    ok = err = nod = 0
    by_name: dict[str, int] = {}
    for r in rows:
        st = str(r.get("status") or "").strip().lower()
        if st == "success":
            ok += 1
        elif st == "error":
            err += 1
        elif st == "no_data":
            nod += 1
        nome = _script_stem(r.get("nome_automacao")) or "?"
        by_name[nome] = by_name.get(nome, 0) + 1
    top = max(by_name.items(), key=lambda x: x[1]) if by_name else ("-", 0)
    return {
        "total": total,
        "success": ok,
        "error": err,
        "no_data": nod,
        "taxa_sucesso": round(100.0 * ok / total, 1) if total else 0.0,
        "taxa_erro": round(100.0 * err / total, 1) if total else 0.0,
        "taxa_sem_dados": round(100.0 * nod / total, 1) if total else 0.0,
        "automacoes_distintas": len(by_name),
        "top_automacao": top[0],
        "top_count": top[1],
    }


def _filter_by_date_range(rows: list[dict[str, Any]], start, end) -> list[dict[str, Any]]:
    a, b = start.isoformat(), end.isoformat()
    return [r for r in rows if a <= str(r.get("data") or "")[:10] <= b]


def _period_compare(curr_rows, prev_rows, curr_label: str, prev_label: str) -> dict[str, Any]:
    c, p = _rows_kpi(curr_rows), _rows_kpi(prev_rows)
    out: dict[str, Any] = {"curr_label": curr_label, "prev_label": prev_label, "curr": c, "prev": p}
    for key in ("total", "taxa_sucesso", "taxa_erro", "taxa_sem_dados"):
        diff = float(c[key]) - float(p[key])
        pct = None if not p[key] else round(100.0 * diff / float(p[key]), 1)
        out[f"delta_{key}"] = round(diff, 1)
        out[f"delta_pct_{key}"] = pct
    return out


def _week_windows(today):
    this_start = today - timedelta(days=today.weekday())
    prev_start = this_start - timedelta(days=7)
    prev_end = this_start - timedelta(days=1)
    return (this_start, today), (prev_start, prev_end)


def _month_windows(today):
    curr_start = today.replace(day=1)
    prev_end = curr_start - timedelta(days=1)
    prev_start = prev_end.replace(day=1)
    return (curr_start, today), (prev_start, prev_end)


def report_data(
    *, year: Optional[int] = None, month: Optional[int] = None, day: Optional[int] = None
) -> dict[str, Any]:
    rows = _read_history_rows_no_catchup()
    today = now_local().date()

    years = sorted({y for r in rows if (y := _to_int_or_none(r.get("ano"))) is not None}, reverse=True)
    months: list[int] = []
    days: list[int] = []
    if year is not None:
        months = sorted(
            {
                m
                for r in rows
                if _to_int_or_none(r.get("ano")) == year and (m := _to_int_or_none(r.get("mes"))) is not None
            }
        )
    if year is not None and month is not None:
        days = sorted(
            {
                d
                for r in rows
                if _to_int_or_none(r.get("ano")) == year
                and _to_int_or_none(r.get("mes")) == month
                and (d := _to_int_or_none(r.get("dia"))) is not None
            }
        )

    scoped = rows
    label_parts: list[str] = []
    if year is not None:
        scoped = [r for r in scoped if _to_int_or_none(r.get("ano")) == year]
        label_parts.append(str(year))
    if month is not None:
        scoped = [r for r in scoped if _to_int_or_none(r.get("mes")) == month]
        label_parts.insert(0, f"{month:02d}")
    if day is not None:
        scoped = [r for r in scoped if _to_int_or_none(r.get("dia")) == day]
        label_parts.insert(0, f"{day:02d}")
    period_label = "/".join(label_parts) if label_parts else "Todo o período"

    trend: list[dict[str, Any]]
    if day is not None:
        trend = []
        granularity = "none"
    elif month is not None and year is not None:
        last_day = calendar.monthrange(year, month)[1]
        bucket = {d: {"success": 0, "error": 0, "no_data": 0, "total": 0} for d in range(1, last_day + 1)}
        for r in scoped:
            d = _to_int_or_none(r.get("dia"))
            if d not in bucket:
                continue
            st = str(r.get("status") or "").strip().lower()
            if st in bucket[d]:
                bucket[d][st] += 1
            bucket[d]["total"] += 1
        trend = [{"label": f"{d:02d}", **v} for d, v in sorted(bucket.items())]
        granularity = "day"
    elif year is not None:
        bucket = {m: {"success": 0, "error": 0, "no_data": 0, "total": 0} for m in range(1, 13)}
        for r in scoped:
            m = _to_int_or_none(r.get("mes"))
            if m not in bucket:
                continue
            st = str(r.get("status") or "").strip().lower()
            if st in bucket[m]:
                bucket[m][st] += 1
            bucket[m]["total"] += 1
        trend = [{"label": _MESES_PT[m - 1], **v} for m, v in sorted(bucket.items())]
        granularity = "month"
    else:
        bucket = {}
        for i in range(29, -1, -1):
            d = today - timedelta(days=i)
            bucket[d.isoformat()] = {"success": 0, "error": 0, "no_data": 0, "total": 0}
        for r in rows:
            dd = str(r.get("data") or "")[:10]
            if dd not in bucket:
                continue
            st = str(r.get("status") or "").strip().lower()
            if st in bucket[dd]:
                bucket[dd][st] += 1
            bucket[dd]["total"] += 1
        trend = [{"label": k[8:10] + "/" + k[5:7], **v} for k, v in bucket.items()]
        granularity = "day"

    w_curr, w_prev = _week_windows(today)
    m_curr, m_prev = _month_windows(today)
    week_block = _period_compare(
        _filter_by_date_range(rows, *w_curr),
        _filter_by_date_range(rows, *w_prev),
        f"{w_curr[0].strftime('%d/%m')} – {w_curr[1].strftime('%d/%m')}",
        f"{w_prev[0].strftime('%d/%m')} – {w_prev[1].strftime('%d/%m')}",
    )
    month_block = _period_compare(
        _filter_by_date_range(rows, *m_curr),
        _filter_by_date_range(rows, *m_prev),
        f"{m_curr[0].strftime('%d/%m')} – {m_curr[1].strftime('%d/%m')}",
        f"{m_prev[0].strftime('%d/%m')} – {m_prev[1].strftime('%d/%m')}",
    )

    return {
        "ok": True,
        "filters": {"year": year, "month": month, "day": day},
        "available": {"years": years, "months": months, "days": days},
        "period_label": period_label,
        "period": _rows_kpi(scoped),
        "trend": trend,
        "trend_granularity": granularity,
        "all_time": _rows_kpi(rows),
        "week": week_block,
        "month_cmp": month_block,
        "total_rows": len(rows),
    }


def _parse_run_dt(raw: str) -> Optional[datetime]:
    s = (raw or "").strip()
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s)
    except ValueError:
        return None
    return _to_sp(dt)


# ---------------------------------------------------------------------------
# Executor + fila (max N em paralelo) + catch-up do dia
# ---------------------------------------------------------------------------

_running: dict[str, dict[str, Any]] = {}
_running_lock = threading.Lock()
_manual_run_last: dict[str, float] = {}
_manual_run_lock = threading.Lock()
_last_fire: dict[str, str] = {}  # nome -> "YYYY-MM-DD HH:MM" already fired by cron tick
_stop_event = threading.Event()

# Fila: mais atrasado primeiro. Dedup por chave nome|scheduled_for|reason
_job_queue: list[dict[str, Any]] = []
_queue_lock = threading.Lock()
_queue_keys: set[str] = set()


def _truncate(text: str, n: int = 8000) -> str:
    text = text or ""
    return text if len(text) <= n else text[: n - 20] + "\n… [truncado]"


def status_from_exit_code(code: int | None) -> str:
    """Convenção das automacoes: 0=success, 1=error, 2=no_data.

    no_data = a automacao terminou sem falha tecnica, mas nao havia material
    para processar (ex.: pasta vazia, email sem anexo naquele momento).
    """
    if code == 0:
        return "success"
    if code == 2:
        return "no_data"
    return "error"


def _queue_key(nome: str, scheduled_for: str, reason: str) -> str:
    return f"{nome.lower()}|{scheduled_for}|{reason}"


def _running_count() -> int:
    with _running_lock:
        return len(_running)


def enqueue_job(
    nome: str,
    path: str,
    reason: str,
    *,
    scheduled_for: Optional[datetime] = None,
    area: Optional[str] = None,
) -> bool:
    """Coloca na fila (dedupe). Retorna True se entrou na fila."""
    now = now_local()
    sched = scheduled_for or now
    if sched.tzinfo is None:
        sched = sched.replace(tzinfo=_tz())
    sched = sched.astimezone(_tz())
    sched_iso = sched.isoformat(timespec="seconds")
    overdue = int((now - sched).total_seconds())  # negativo = ainda no futuro
    key = _queue_key(nome, sched_iso, reason)
    with _queue_lock:
        if key in _queue_keys:
            return False
        _queue_keys.add(key)
        _job_queue.append(
            {
                "key": key,
                "nome_automacao": nome,
                "path": path,
                "area": area,
                "reason": reason,
                "scheduled_for": sched_iso,
                "overdue_sec": max(0, overdue),
                "wait_sec": max(0, -overdue),
                "due": overdue >= 0 or reason == "manual",
                "enqueued_at": now_iso(),
            }
        )
        _sort_queue_locked()
    log.info(
        "[QUEUE] + %s reason=%s scheduled_for=%s due=%s (fila=%d)",
        nome,
        reason,
        sched_iso,
        overdue >= 0 or reason == "manual",
        len(_job_queue),
    )
    return True


def _sort_queue_locked() -> None:
    """Prontas (atrasadas) primeiro; depois as que ainda vao chegar, por horario."""
    now = now_local()

    def sort_key(j: dict[str, Any]) -> tuple:
        try:
            sf = datetime.fromisoformat(j["scheduled_for"])
            if sf.tzinfo is None:
                sf = sf.replace(tzinfo=_tz())
            else:
                sf = sf.astimezone(_tz())
        except ValueError:
            sf = now
        overdue = int((now - sf).total_seconds())
        j["overdue_sec"] = max(0, overdue)
        j["wait_sec"] = max(0, -overdue)
        is_manual = j.get("reason") == "manual"
        due = is_manual or overdue >= 0
        j["due"] = due
        # due first (0), then higher overdue, then earlier scheduled_for
        return (0 if due else 1, -max(0, overdue) if due else 0, j["scheduled_for"])

    _job_queue.sort(key=sort_key)


def list_queue() -> list[dict[str, Any]]:
    with _queue_lock:
        _sort_queue_locked()
        return [dict(j) for j in _job_queue]


def _pop_next_runnable() -> Optional[dict[str, Any]]:
    """Tira da fila a mais atrasada que JA PODE correr (horario chegou) e nao esta busy.

    Reserva o nome em ``_running`` imediatamente (antes da thread) para nunca
    disparar a mesma automacao em paralelo.
    """
    now = now_local()
    with _queue_lock:
        _sort_queue_locked()
        with _running_lock:
            busy = set(_running.keys())
            for i, job in enumerate(_job_queue):
                nome = job["nome_automacao"]
                if nome in busy:
                    continue
                reason = job.get("reason") or ""
                try:
                    sf = datetime.fromisoformat(job["scheduled_for"])
                    if sf.tzinfo is None:
                        sf = sf.replace(tzinfo=_tz())
                    else:
                        sf = sf.astimezone(_tz())
                except ValueError:
                    sf = now
                if reason != "manual" and sf > now:
                    continue
                if len(_running) >= MAX_CONCURRENT:
                    return None
                item = _job_queue.pop(i)
                _queue_keys.discard(item["key"])
                # reserva ja: impede segundo pop da mesma automacao
                _running[nome] = {
                    "started": now_iso(),
                    "pid": None,
                    "reason": item.get("reason"),
                    "reserved": True,
                }
                busy.add(nome)
                return item
    return None


def run_script(
    nome: str,
    path: str,
    reason: str = "manual",
    *,
    pre_reserved: bool = False,
    area: Optional[str] = None,
) -> dict[str, Any]:
    with _running_lock:
        if pre_reserved:
            info = _running.get(nome)
            if not info:
                _running[nome] = {"started": now_iso(), "pid": None, "reason": reason}
            else:
                info["reason"] = reason
                info.pop("reserved", None)
        else:
            if nome in _running:
                return {"ok": False, "error": "ja em execucao", "nome_automacao": nome}
            if len(_running) >= MAX_CONCURRENT:
                return {"ok": False, "error": f"limite de {MAX_CONCURRENT} em paralelo", "nome_automacao": nome}
            _running[nome] = {"started": now_iso(), "pid": None, "reason": reason}

    if not area:
        _, area_found = find_script(nome)
        area = area_found or ""

    start = time.time()
    start_iso = now_iso()
    rid = insert_run(
        {
            "nome_automacao": nome,
            "area": area or "",
            "trigger_reason": reason,
            "status": "running",
            "start_time": start_iso,
        }
    )
    log.info("[RUN] %s (%s) path=%s concurrent=%d/%d", nome, reason, path, _running_count(), MAX_CONCURRENT)
    try:
        proc = subprocess.Popen(
            [sys.executable, path],
            cwd=str(Path(path).parent),
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        with _running_lock:
            if nome in _running:
                _running[nome]["pid"] = proc.pid
        try:
            out, _ = proc.communicate(timeout=JOB_TIMEOUT_SEC or None)
            code = proc.returncode if proc.returncode is not None else 1
            status = status_from_exit_code(code)
        except subprocess.TimeoutExpired:
            _kill_pid(proc.pid)
            try:
                out, _ = proc.communicate(timeout=10)
            except subprocess.TimeoutExpired:
                out = ""
            code = -1
            status = "error"
            out = (out or "") + f"\n[ServerCRON] Timeout: automacao excedeu {JOB_TIMEOUT_SEC}s e foi encerrada."
            log.error("[RUN] timeout %s (%ds) — processo encerrado", nome, JOB_TIMEOUT_SEC)
    except Exception as exc:
        out = str(exc)
        code = 1
        status = "error"
        log.exception("[RUN] falha %s", nome)
    finally:
        with _running_lock:
            _running.pop(nome, None)

    duration = round(time.time() - start, 2)
    end_iso = now_iso()
    update_run(
        rid,
        status=status,
        end_time=end_iso,
        duration_sec=duration,
        exit_code=code,
        output=_truncate(out),
    )
    log.info("[DONE] %s status=%s exit=%s %.2fs", nome, status, code, duration)
    return {
        "ok": True,
        "id": rid,
        "nome_automacao": nome,
        "status": status,
        "exit_code": code,
        "duration_sec": duration,
        "start_time": start_iso,
        "end_time": end_iso,
    }


def _kill_pid(pid: int) -> bool:
    try:
        if sys.platform == "win32":
            subprocess.call(
                ["taskkill", "/PID", str(pid), "/T", "/F"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        else:
            os.kill(pid, signal.SIGKILL)
        return True
    except OSError:
        return False


def kill_script(nome: str) -> bool:
    with _running_lock:
        info = _running.get(nome)
        pid = info.get("pid") if info else None
    if not pid:
        return False
    return _kill_pid(pid)


def _minute_key(dt: datetime) -> str:
    return dt.strftime("%Y-%m-%d %H:%M")


def _should_fire(cron_expr: str, moment: datetime) -> bool:
    """True se este minuto (horario de Brasilia) e um tick do cron."""
    try:
        moment_sp = _to_sp(moment).replace(second=0, microsecond=0)
        prev = _croniter_sp(cron_expr, moment_sp).get_prev(datetime)
        prev = _to_sp(prev)
        return _minute_key(prev) == _minute_key(moment_sp)
    except Exception:
        return False


def _cron_fire_times_today(cron_expr: str, until: datetime) -> list[datetime]:
    """Todos os ticks do cron desde 00:00 do dia (Brasilia) ate ``until``."""
    until_sp = _to_sp(until).replace(second=0, microsecond=0)
    start = until_sp.replace(hour=0, minute=0, second=0, microsecond=0)
    cursor = start - timedelta(seconds=1)
    times: list[datetime] = []
    try:
        it = _croniter_sp(cron_expr, cursor)
        for _ in range(2000):
            nxt = _to_sp(it.get_next(datetime))
            if nxt.date() != until_sp.date() or nxt > until_sp:
                break
            times.append(nxt)
    except Exception:
        return []
    return times


def _cron_fire_times_rest_of_day(cron_expr: str, after: datetime) -> list[datetime]:
    """Ticks do cron estritamente apos ``after`` ate ao fim do dia (Brasilia)."""
    after_sp = _to_sp(after).replace(second=0, microsecond=0)
    end = after_sp.replace(hour=23, minute=59, second=59, microsecond=0)
    times: list[datetime] = []
    try:
        it = _croniter_sp(cron_expr, after_sp)
        for _ in range(2000):
            nxt = _to_sp(it.get_next(datetime))
            if nxt.date() != after_sp.date() or nxt > end:
                break
            times.append(nxt)
    except Exception:
        return []
    return times


def detect_pending_today() -> list[dict[str, Any]]:
    """Automacoes ativas cujo(s) horario(s) de hoje ja passaram e ainda nao correram.

    Matching guloso: cada execucao do dia cobre o horario pendente mais antigo
    com fire <= start (assim um catch-up as 15h cobre o slot das 07h).
    Ordenado do mais atrasado para o menos.
    """
    now = now_local().replace(second=0, microsecond=0)
    history = list_history_today()
    runs_by_name: dict[str, list[datetime]] = {}
    for h in history:
        nome = _script_stem(h.get("nome_automacao"))
        if not nome:
            continue
        dt = _parse_run_dt(str(h.get("start_time") or ""))
        if not dt:
            continue
        runs_by_name.setdefault(nome.lower(), []).append(dt)

    pending: list[dict[str, Any]] = []
    for job in schedulable(force=False):
        nome = job["nome_automacao"]
        cron_expr = job["cron_schedule"]
        path = job["path"]
        if not path:
            continue
        fires = _cron_fire_times_today(cron_expr, now)
        if not fires:
            continue
        runs = sorted(runs_by_name.get(nome.lower(), []))
        used = [False] * len(runs)
        missed: list[datetime] = []
        for fire in fires:
            covered = False
            for i, rs in enumerate(runs):
                if used[i]:
                    continue
                if rs >= fire:
                    used[i] = True
                    covered = True
                    break
            if not covered:
                missed.append(fire)
        missed.sort()
        for fire in missed[:CATCHUP_PER_SCRIPT]:
            overdue = max(0, int((now - fire).total_seconds()))
            pending.append(
                {
                    "nome_automacao": nome,
                    "area": job.get("area"),
                    "path": path,
                    "cron_schedule": cron_expr,
                    "scheduled_for": fire.isoformat(timespec="seconds"),
                    "overdue_sec": overdue,
                    "overdue_human": _format_overdue(overdue),
                }
            )

    pending.sort(key=lambda p: (-int(p["overdue_sec"]), p["scheduled_for"]))
    return pending


def _format_overdue(seconds: int) -> str:
    if seconds < 60:
        return f"{seconds}s"
    if seconds < 3600:
        return f"{seconds // 60} min"
    h = seconds // 3600
    m = (seconds % 3600) // 60
    if m:
        return f"{h}h {m}min"
    return f"{h}h"


def list_upcoming_today() -> list[dict[str, Any]]:
    """Proximos horarios de hoje (ainda por vir) por automacao ativa — fuso Brasilia."""
    now = now_local().replace(second=0, microsecond=0)
    upcoming: list[dict[str, Any]] = []
    for job in schedulable(force=False):
        path = job["path"]
        if not path:
            continue
        fires = _cron_fire_times_rest_of_day(job["cron_schedule"], now)[:UPCOMING_PER_SCRIPT]
        for fire in fires:
            wait = max(0, int((fire - now).total_seconds()))
            upcoming.append(
                {
                    "nome_automacao": job["nome_automacao"],
                    "area": job.get("area"),
                    "path": path,
                    "cron_schedule": job["cron_schedule"],
                    "scheduled_for": fire.isoformat(timespec="seconds"),
                    "wait_sec": wait,
                    "wait_human": _format_overdue(wait),
                }
            )
    upcoming.sort(key=lambda u: u["scheduled_for"])
    return upcoming


def enqueue_catchup(pending: Optional[list[dict[str, Any]]] = None) -> int:
    """Enfileira pendentes do dia (mais atrasados primeiro)."""
    items = pending if pending is not None else detect_pending_today()
    n = 0
    for p in items:
        try:
            sched = datetime.fromisoformat(p["scheduled_for"])
        except ValueError:
            continue
        if enqueue_job(
            p["nome_automacao"],
            p["path"],
            "catchup",
            scheduled_for=sched,
            area=p.get("area"),
        ):
            n += 1
    return n


def enqueue_upcoming(upcoming: Optional[list[dict[str, Any]]] = None) -> int:
    """Enfileira proximos horarios do dia; so correm quando a hora chegar."""
    items = upcoming if upcoming is not None else list_upcoming_today()
    n = 0
    for u in items:
        try:
            sched = datetime.fromisoformat(u["scheduled_for"])
        except ValueError:
            continue
        if enqueue_job(
            u["nome_automacao"],
            u["path"],
            "scheduled",
            scheduled_for=sched,
            area=u.get("area"),
        ):
            n += 1
    return n


def refresh_day_queue() -> dict[str, int]:
    """Atrasadas + proximas do dia na mesma fila."""
    c = enqueue_catchup()
    u = enqueue_upcoming()
    return {"catchup_enqueued": c, "upcoming_enqueued": u, "queue_count": len(list_queue())}


def clear_schedule_queue(*, keep_manual: bool = True) -> int:
    """Limpa jobs de catchup/scheduled da fila (ex.: apos mudanca na planilha)."""
    removed = 0
    with _queue_lock:
        keep: list[dict[str, Any]] = []
        for j in _job_queue:
            if keep_manual and j.get("reason") == "manual":
                keep.append(j)
                continue
            _queue_keys.discard(j["key"])
            removed += 1
        _job_queue[:] = keep
    return removed


def reload_spreadsheet_and_schedules() -> dict[str, Any]:
    """Relê a planilha e reconstrói a fila de agendamentos (tempo real)."""
    scripts = read_registry(force=True)
    inv = inventory_summary(scripts)
    cleared = clear_schedule_queue(keep_manual=True)
    # permite re-disparar o minuto atual se o cron mudou
    _last_fire.clear()
    stats = refresh_day_queue()
    mark_spreadsheet_reload()
    pending = detect_pending_today()
    queue = list_queue()
    log.info(
        "[RELOAD] planilha=%d scripts | fila limpa=%d | catchup=+%d upcoming=+%d | fila=%d",
        len(scripts),
        cleared,
        stats["catchup_enqueued"],
        stats["upcoming_enqueued"],
        stats["queue_count"],
    )
    meta = spreadsheet_reload_meta()
    return {
        "scripts": scripts,
        "inventory": inv,
        "count": len(scripts),
        "queue_cleared": cleared,
        "catchup_enqueued": stats["catchup_enqueued"],
        "upcoming_enqueued": stats["upcoming_enqueued"],
        "queue_count": stats["queue_count"],
        "pending_count": len(pending),
        "pending": pending,
        "queue": queue,
        "queue_due": len([j for j in queue if j.get("due")]),
        "queue_waiting": len([j for j in queue if not j.get("due")]),
        **meta,
    }


def queue_processor_loop() -> None:
    """Consome a fila respeitando MAX_CONCURRENT (default 5)."""
    log.info("[QUEUE] processador iniciado | max_concurrent=%d", MAX_CONCURRENT)
    while not _stop_event.is_set():
        try:
            while _running_count() < MAX_CONCURRENT:
                job = _pop_next_runnable()
                if not job:
                    break
                nome = job["nome_automacao"]
                path = job["path"]
                reason = job["reason"]
                # marca scheduled fire para nao duplicar no tick
                if reason in ("scheduled", "catchup") and job.get("scheduled_for"):
                    try:
                        sf = datetime.fromisoformat(job["scheduled_for"])
                        _last_fire[nome] = _minute_key(sf)
                    except ValueError:
                        pass

                def _bg(n=nome, p=path, r=reason, a=job.get("area")):
                    try:
                        run_script(n, p, r, pre_reserved=True, area=a)
                    except Exception:
                        log.exception("[QUEUE] falha ao correr %s", n)
                        with _running_lock:
                            _running.pop(n, None)

                threading.Thread(target=_bg, daemon=True, name=f"job-{nome}").start()
                log.info(
                    "[QUEUE] disparou %s (%s) atrasado=%ss concurrent=%d/%d",
                    nome,
                    reason,
                    job.get("overdue_sec"),
                    _running_count(),
                    MAX_CONCURRENT,
                )
        except Exception:
            log.exception("[QUEUE] erro no processador")
        _stop_event.wait(2)


def scheduler_loop() -> None:
    log.info(
        "[SCHED] loop iniciado | reload %d min | catchup+upcoming | tz=%s | max_concurrent=%d",
        RELOAD_MINUTES,
        TZ_NAME,
        MAX_CONCURRENT,
    )
    last_queue_refresh = 0.0
    # catch-up + proximas do dia ao arrancar (e marca leitura inicial da planilha)
    try:
        read_registry(force=True)
        stats = refresh_day_queue()
        mark_spreadsheet_reload()
        log.info(
            "[QUEUE] boot: catchup=+%d upcoming=+%d fila=%d | proximo reload em %d min",
            stats["catchup_enqueued"],
            stats["upcoming_enqueued"],
            stats["queue_count"],
            RELOAD_MINUTES,
        )
    except Exception:
        log.exception("[QUEUE] boot falhou")
        mark_spreadsheet_reload()

    while not _stop_event.is_set():
        try:
            now = now_local().replace(second=0, microsecond=0)
            # a cada RELOAD_MINUTES: rele planilha e reconstrói fila de cron
            with _spreadsheet_reload_lock:
                last_ts = float(_last_spreadsheet_reload_ts or 0.0)
            if last_ts <= 0 or time.time() - last_ts >= RELOAD_MINUTES * 60:
                try:
                    reload_spreadsheet_and_schedules()
                except Exception:
                    log.exception("[RELOAD] falha no auto-reload da planilha")
                    mark_spreadsheet_reload()
                last_queue_refresh = time.time()

            # reavalia atrasadas + proximas a cada 2 min (entre reloads)
            if time.time() - last_queue_refresh >= 120:
                stats = refresh_day_queue()
                if stats["catchup_enqueued"] or stats["upcoming_enqueued"]:
                    log.info(
                        "[QUEUE] refresh: catchup=+%d upcoming=+%d fila=%d",
                        stats["catchup_enqueued"],
                        stats["upcoming_enqueued"],
                        stats["queue_count"],
                    )
                last_queue_refresh = time.time()

            # tick do minuto atual (garantia alem do pre-agendamento)
            for job in schedulable(force=False):
                nome = job["nome_automacao"]
                cron_expr = job["cron_schedule"]
                path = job["path"]
                if not path:
                    continue
                if not _should_fire(cron_expr, now):
                    continue
                if _last_fire.get(nome) == _minute_key(now):
                    continue
                enqueue_job(
                    nome,
                    path,
                    "scheduled",
                    scheduled_for=now,
                    area=job.get("area"),
                )
        except Exception:
            log.exception("[SCHED] erro no loop")
        _stop_event.wait(15)


# ---------------------------------------------------------------------------
# Clima (Open-Meteo — open source, Sao Paulo)
# ---------------------------------------------------------------------------

_weather_lock = threading.Lock()
_weather_cache: dict[str, Any] = {"ts": 0.0, "data": None, "error": None}


def _wmo_label_pt(code: int) -> str:
    """Traducao resumida dos weather codes WMO (Open-Meteo)."""
    c = int(code)
    if c == 0:
        return "Céu limpo"
    if c in (1, 2):
        return "Parcialmente nublado"
    if c == 3:
        return "Nublado"
    if c in (45, 48):
        return "Nevoeiro"
    if c in (51, 53, 55, 56, 57):
        return "Garoa"
    if c in (61, 63, 65, 66, 67):
        return "Chuva"
    if c in (71, 73, 75, 77):
        return "Neve"
    if c in (80, 81, 82):
        return "Aguaceiros"
    if c in (85, 86):
        return "Aguaceiros de neve"
    if c == 95:
        return "Trovoada"
    if c in (96, 99):
        return "Trovoada com granizo"
    return "Indefinido"


def _wmo_icon(code: int) -> str:
    c = int(code)
    if c == 0:
        return "clear"
    if c in (1, 2, 3):
        return "cloudy"
    if c in (45, 48):
        return "fog"
    if c in (71, 73, 75, 77, 85, 86):
        return "snow"
    if c in (95, 96, 99):
        return "storm"
    if c in (51, 53, 55, 56, 57, 61, 63, 65, 66, 67, 80, 81, 82):
        return "rain"
    return "cloudy"


def _fetch_open_meteo_sp() -> dict[str, Any]:
    qs = urllib.parse.urlencode(
        {
            "latitude": WEATHER_LAT,
            "longitude": WEATHER_LON,
            "current": "temperature_2m,apparent_temperature,relative_humidity_2m,weather_code,wind_speed_10m",
            "daily": "weather_code,temperature_2m_max,temperature_2m_min",
            "timezone": "America/Sao_Paulo",
            "forecast_days": 7,
        }
    )
    url = f"https://api.open-meteo.com/v1/forecast?{qs}"
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "ServerCRON/1.0 (+https://github.com/caducosilva/ServerCRON; open-meteo)"},
        method="GET",
    )
    with urllib.request.urlopen(req, timeout=12) as resp:
        raw = json.loads(resp.read().decode("utf-8"))

    cur = raw.get("current") or {}
    daily = raw.get("daily") or {}
    code = int(cur.get("weather_code") if cur.get("weather_code") is not None else 0)
    temp = cur.get("temperature_2m")
    tmax = (daily.get("temperature_2m_max") or [None])[0]
    tmin = (daily.get("temperature_2m_min") or [None])[0]

    days: list[dict[str, Any]] = []
    times = daily.get("time") or []
    codes = daily.get("weather_code") or []
    maxs = daily.get("temperature_2m_max") or []
    mins = daily.get("temperature_2m_min") or []
    for i, day in enumerate(times):
        if len(days) >= 7:
            break
        dcode = int(codes[i]) if i < len(codes) and codes[i] is not None else 0
        is_today = i == 0
        days.append(
            {
                "date": day,
                "is_today": is_today,
                "weather_code": code if is_today else dcode,
                "label": _wmo_label_pt(code if is_today else dcode),
                "icon": _wmo_icon(code if is_today else dcode),
                "temperature": temp if is_today else (maxs[i] if i < len(maxs) else None),
                "temp_max": maxs[i] if i < len(maxs) else None,
                "temp_min": mins[i] if i < len(mins) else None,
            }
        )

    return {
        "ok": True,
        "source": "open-meteo",
        "license": "CC BY 4.0 (Open-Meteo)",
        "city": "São Paulo",
        "country": "Brasil",
        "timezone": "America/Sao_Paulo",
        "latitude": WEATHER_LAT,
        "longitude": WEATHER_LON,
        "temperature": temp,
        "apparent_temperature": cur.get("apparent_temperature"),
        "humidity": cur.get("relative_humidity_2m"),
        "wind_speed": cur.get("wind_speed_10m"),
        "weather_code": code,
        "label": _wmo_label_pt(code),
        "icon": _wmo_icon(code),
        "temp_max": tmax,
        "temp_min": tmin,
        "observed_at": cur.get("time"),
        "days": days,
        "units": (raw.get("current_units") or {}),
    }


def get_weather_sp(*, force: bool = False) -> dict[str, Any]:
    """Clima atual de Sao Paulo com cache em memoria."""
    now = time.time()
    with _weather_lock:
        cached = _weather_cache.get("data")
        if (
            not force
            and cached
            and now - float(_weather_cache.get("ts") or 0) < WEATHER_CACHE_SEC
        ):
            out = dict(cached)
            out["cached"] = True
            out["cache_age_sec"] = int(now - float(_weather_cache["ts"]))
            return out

    try:
        data = _fetch_open_meteo_sp()
        with _weather_lock:
            _weather_cache["ts"] = time.time()
            _weather_cache["data"] = data
            _weather_cache["error"] = None
        data = dict(data)
        data["cached"] = False
        data["cache_age_sec"] = 0
        return data
    except Exception as exc:
        log.warning("[WEATHER] falha Open-Meteo: %s", exc)
        with _weather_lock:
            stale = _weather_cache.get("data")
            _weather_cache["error"] = str(exc)
            if stale:
                out = dict(stale)
                out["cached"] = True
                out["stale"] = True
                out["error"] = str(exc)
                out["cache_age_sec"] = int(now - float(_weather_cache.get("ts") or now))
                return out
        return {
            "ok": False,
            "source": "open-meteo",
            "city": "São Paulo",
            "country": "Brasil",
            "error": str(exc),
            "days": [],
        }


# ---------------------------------------------------------------------------
# Flask app
# ---------------------------------------------------------------------------

app = Flask(__name__, static_folder=str(BASE_DIR / "static"), static_url_path="/static")
STARTED_AT = now_iso()


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _request_token() -> str:
    header = request.headers.get("X-ServerCRON-Token") or ""
    if not header:
        auth = request.headers.get("Authorization") or ""
        if auth.lower().startswith("bearer "):
            header = auth[7:]
    return header.strip()


def _auth_configured() -> bool:
    """Protecao ativa se houver token estatico OU pelo menos 1 email cadastrado na aba USERS."""
    if AUTH_DISABLED:
        return False
    return bool(API_TOKEN) or bool(read_users())


def _token_ok() -> bool:
    """Valido por token estatico (tempo constante) OU por sessao de login por email (nao expirada)."""
    if not _auth_configured():
        return True
    token = _request_token()
    if not token:
        return False
    if API_TOKEN and hmac.compare_digest(token, API_TOKEN):
        return True
    return session_email(token) is not None


def _client_ip() -> str:
    return request.remote_addr or "desconhecido"


_PUBLIC_API_PATHS = ("/api/auth", "/api/login/request", "/api/login/verify")


@app.before_request
def _require_api_token():
    if not _auth_configured():
        return None
    if not request.path.startswith("/api/") or request.path in _PUBLIC_API_PATHS:
        return None
    if not _token_ok():
        return jsonify({"ok": False, "error": "sessao invalida ou expirada, faca login novamente"}), 401
    return None


def _require_session_email() -> Optional[str]:
    """P/ logs: quem esta logado (email da sessao, ou 'token-admin' se via SERVERCRON_API_TOKEN)."""
    token = _request_token()
    if API_TOKEN and hmac.compare_digest(token, API_TOKEN):
        return "token-admin"
    return session_email(token)


def _current_email() -> Optional[str]:
    token = _request_token()
    return session_email(token) if token else None


def _current_is_admin() -> bool:
    """SERVERCRON_API_TOKEN (dono do servidor) sempre conta como admin; senao, checa is_admin na USERS."""
    token = _request_token()
    if API_TOKEN and hmac.compare_digest(token, API_TOKEN):
        return True
    email = session_email(token)
    return bool(email) and _email_is_admin(email)


def _login_mode() -> str:
    """otp = aba USERS configurada -> login exige codigo enviado por email (unico fator que
    realmente prova posse do email; token estatico sozinho nao prova nada).
    token = modo legado, sem aba USERS, so SERVERCRON_API_TOKEN."""
    if read_users():
        return "otp"
    if API_TOKEN:
        return "token"
    return "token"


@app.get("/api/auth")
def api_auth():
    """Publico: diz ao painel se precisa de login e por qual metodo (sem revelar segredos)."""
    token = _request_token()
    email = session_email(token) if token else None
    return jsonify(
        {
            "ok": True,
            "auth_required": _auth_configured(),
            "authenticated": _token_ok(),
            "login_mode": _login_mode(),
            "email": email,
            "is_admin": _current_is_admin(),
            "session_ttl_sec": SESSION_TTL_SEC,
            "otp_ttl_sec": OTP_TTL_SEC,
        }
    )


@app.post("/api/login/request")
def api_login_request():
    """Passo 1: se o email estiver cadastrado e ativo, gera um codigo aleatorio de uso unico
    e envia por email (SMTP). Resposta e sempre a mesma independente do email existir ou nao,
    pra nao dar pra descobrir por aqui quem esta cadastrado (so quem tem a caixa de entrada
    sabe se recebeu algo)."""
    ip = _client_ip()
    if _rate_limited("otp-req-ip", ip, max_hits=8, window_sec=300):
        return jsonify({"ok": False, "error": "muitas tentativas deste IP, aguarde alguns minutos"}), 429
    data = request.get_json(silent=True) or {}
    email = _email_norm(data.get("email"))
    if not email or not _EMAIL_RE.match(email):
        return jsonify({"ok": False, "error": "email invalido"}), 400
    if _rate_limited("otp-req-email", email, max_hits=5, window_sec=900):
        return jsonify({"ok": False, "error": "muitos pedidos de codigo pra este email, aguarde alguns minutos"}), 429

    generic_response = jsonify(
        {"ok": True, "message": "Se o email estiver cadastrado, um codigo foi enviado.", "ttl_sec": OTP_TTL_SEC}
    )

    if not _email_allowed(email):
        log.warning("[OTP] pedido de codigo para email nao cadastrado/inativo %s (ip=%s)", email, ip)
        return generic_response

    if not _smtp_configured():
        log.error("[OTP] SERVERCRON_SMTP_USER/SERVERCRON_SMTP_PASSWORD nao configurados — nao consigo enviar codigo")
        return jsonify({"ok": False, "error": "envio de email nao configurado no servidor"}), 500

    code = _generate_otp()
    with _otp_lock:
        _otp_codes[email] = {"code": code, "expires_at": time.time() + OTP_TTL_SEC, "attempts": 0}
    try:
        _send_otp_email(email, code)
    except Exception:
        log.exception("[OTP] falha ao enviar email para %s", email)
        with _otp_lock:
            _otp_codes.pop(email, None)
        return jsonify({"ok": False, "error": "falha ao enviar o email, tente novamente"}), 502

    log.info("[OTP] codigo enviado para %s (ip=%s)", email, ip)
    return generic_response


@app.post("/api/login/verify")
def api_login_verify():
    """Passo 2: confirma o codigo recebido por email. So AQUI uma sessao e emitida —
    e so aqui uma sessao anterior do mesmo email e derrubada (pedir codigo sozinho nunca
    desloga ninguem; e preciso provar que recebeu o email)."""
    ip = _client_ip()
    if _login_rate_limited(ip):
        return jsonify({"ok": False, "error": "muitas tentativas, aguarde alguns minutos"}), 429
    data = request.get_json(silent=True) or {}
    email = _email_norm(data.get("email"))
    code_in = str(data.get("code") or "").strip()
    _register_login_attempt(ip)
    if not email or not code_in:
        return jsonify({"ok": False, "error": "informe email e codigo"}), 400
    if _rate_limited("otp-verify-email", email, max_hits=OTP_MAX_VERIFY_ATTEMPTS + 2, window_sec=OTP_TTL_SEC):
        return jsonify({"ok": False, "error": "muitas tentativas pra este email, peca um novo codigo"}), 429

    with _otp_lock:
        entry = _otp_codes.get(email)
        valid = False
        if entry and entry["expires_at"] > time.time() and entry["attempts"] < OTP_MAX_VERIFY_ATTEMPTS:
            entry["attempts"] += 1
            valid = hmac.compare_digest(code_in, entry["code"])
        if entry and (valid or entry["expires_at"] <= time.time() or entry["attempts"] >= OTP_MAX_VERIFY_ATTEMPTS):
            del _otp_codes[email]

    if not valid or not _email_allowed(email):
        log.warning("[OTP] verificacao falhou para %s (ip=%s)", email, ip)
        return jsonify({"ok": False, "error": "codigo invalido ou expirado"}), 403

    session = issue_session(email)
    log.info("[LOGIN] sessao criada via codigo por email para %s (ip=%s)", email, ip)
    return jsonify(
        {
            "ok": True,
            "token": session["token"],
            "expires_at": session["expires_at"],
            "email": email,
            "is_admin": _email_is_admin(email),
        }
    )


@app.post("/api/logout")
def api_logout():
    token = _request_token()
    if token:
        revoke_session(token)
    return jsonify({"ok": True})


def _require_admin():
    """Retorna uma resposta 403 se quem chamou nao for admin; None se pode prosseguir."""
    if not _current_is_admin():
        who = _current_email() or _client_ip()
        log.warning("[USERS] acesso negado (nao-admin) para %s em %s", who, request.path)
        return jsonify({"ok": False, "error": "apenas administradores podem gerir usuarios"}), 403
    return None


@app.get("/api/users")
def api_users_list():
    denied = _require_admin()
    if denied:
        return denied
    return jsonify({"ok": True, "users": read_users()})


@app.post("/api/users")
def api_users_add():
    denied = _require_admin()
    if denied:
        return denied
    data = request.get_json(silent=True) or {}
    email = _email_norm(data.get("email"))
    if not email or not _EMAIL_RE.match(email):
        return jsonify({"ok": False, "error": "email invalido"}), 400
    nome = _safe_name(data.get("nome")) or None
    users = read_users(force=True)
    if any(u["email"] == email for u in users):
        return jsonify({"ok": False, "error": "email ja cadastrado"}), 409
    users.append({"email": email, "nome": nome, "is_active": True, "is_admin": False})
    write_users(users)
    log.info("[USERS] %s adicionado por %s", email, _require_session_email() or "?")
    return jsonify({"ok": True, "users": read_users(force=True)})


@app.post("/api/users/<path:email>/toggle")
def api_users_toggle(email: str):
    denied = _require_admin()
    if denied:
        return denied
    target = _email_norm(email)
    users = read_users(force=True)
    found = False
    for u in users:
        if u["email"] == target:
            u["is_active"] = not u["is_active"]
            found = True
            break
    if not found:
        return jsonify({"ok": False, "error": "email nao encontrado"}), 404
    write_users(users)
    log.info("[USERS] %s alternado por %s", target, _require_session_email() or "?")
    return jsonify({"ok": True, "users": read_users(force=True)})


@app.delete("/api/users/<path:email>")
def api_users_delete(email: str):
    denied = _require_admin()
    if denied:
        return denied
    target = _email_norm(email)
    users = read_users(force=True)
    remaining = [u for u in users if u["email"] != target]
    if len(remaining) == len(users):
        return jsonify({"ok": False, "error": "email nao encontrado"}), 404
    write_users(remaining)
    log.info("[USERS] %s removido por %s", target, _require_session_email() or "?")
    return jsonify({"ok": True, "users": read_users(force=True)})


@app.after_request
def _security_headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "SAMEORIGIN"
    resp.headers["Referrer-Policy"] = "same-origin"
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; "
        "script-src 'self'; "
        "style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data:; "
        "connect-src 'self'; "
        "frame-ancestors 'self'"
    )
    return resp


@app.errorhandler(Exception)
def _handle_unexpected_error(exc):
    from werkzeug.exceptions import HTTPException

    if isinstance(exc, HTTPException):
        return exc
    log.exception("[HTTP] erro nao tratado em %s", request.path)
    return jsonify({"ok": False, "error": "erro interno do servidor"}), 500


@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.get("/login")
def login_page():
    return send_from_directory(app.static_folder, "login.html")


@app.get("/api/status")
def api_status():
    scripts = read_registry()
    inv = inventory_summary(scripts)
    pending = detect_pending_today()
    with _running_lock:
        live = [
            {
                "nome_automacao": k,
                "started": v.get("started"),
                "pid": v.get("pid"),
                "reason": v.get("reason"),
            }
            for k, v in _running.items()
        ]
    queue = list_queue()
    return jsonify(
        {
            "ok": True,
            "started_at": STARTED_AT,
            "now": now_iso(),
            "tz": TZ_NAME,
            "tz_label": "Horário de Brasília (São Paulo)",
            "now_brasilia": now_iso(),
            "host": HOST,
            "port": PORT,
            "lan_urls": [f"http://{ip}:{PORT}/" for ip in _local_ips()] if HOST not in _LOOPBACK_HOSTS else [],
            "automacoes_dir": str(AUTOMAOES_DIR),
            "areas": list_area_dirs(),
            "data_root": str(DATA_ROOT),
            "home": str(HOME),
            "config_path": str(CONFIG_PATH),
            "registro_xlsx": str(REGISTRO_XLSX),
            "sqlite_path": str(SQLITE_PATH),
            "logs_dir": str(LOGS_DIR),
            "history_csv": str(HISTORY_CSV_PATH),
            "registro_exists": REGISTRO_XLSX.is_file(),
            "automacoes_locked_by_env": _automacoes_from_env,
            "total_scripts": inv["sheet_total"],
            "active_scripts": inv["sheet_active"],
            "inactive_scripts": inv["sheet_inactive"],
            "located_scripts": inv["sheet_located"],
            "missing_scripts": inv["sheet_missing"],
            "py_on_disk": inv["py_on_disk"],
            "areas_count": inv["areas_count"],
            "schedulable": inv["schedulable"],
            "inventory": inv,
            "running": live,
            "running_count": len(live),
            "max_concurrent": MAX_CONCURRENT,
            "pending": pending,
            "pending_count": len(pending),
            "queue": queue,
            "queue_count": len(queue),
            "queue_due": len([j for j in queue if j.get("due")]),
            "queue_waiting": len([j for j in queue if not j.get("due")]),
            **spreadsheet_reload_meta(),
        }
    )


@app.get("/api/pending")
def api_pending():
    pending = detect_pending_today()
    return jsonify(
        {
            "ok": True,
            "pending": pending,
            "count": len(pending),
            "queue": list_queue(),
            "running_count": _running_count(),
            "max_concurrent": MAX_CONCURRENT,
        }
    )


@app.post("/api/catchup")
def api_catchup():
    if _rate_limited("catchup", _client_ip(), max_hits=10, window_sec=60):
        return jsonify({"ok": False, "error": "muitas requisicoes, aguarde um pouco"}), 429
    pending = detect_pending_today()
    upcoming = list_upcoming_today()
    stats = refresh_day_queue()
    queue = list_queue()
    due = [j for j in queue if j.get("due")]
    waiting = [j for j in queue if not j.get("due")]
    return jsonify(
        {
            "ok": True,
            "enqueued": stats["catchup_enqueued"] + stats["upcoming_enqueued"],
            "catchup_enqueued": stats["catchup_enqueued"],
            "upcoming_enqueued": stats["upcoming_enqueued"],
            "pending": pending,
            "pending_count": len(pending),
            "upcoming": upcoming[:40],
            "upcoming_count": len(upcoming),
            "queue": queue,
            "queue_count": len(queue),
            "queue_due": len(due),
            "queue_waiting": len(waiting),
            "max_concurrent": MAX_CONCURRENT,
        }
    )


@app.get("/api/config")
def api_config_get():
    return jsonify(
        {
            "ok": True,
            "automacoes_dir": str(AUTOMAOES_DIR),
            "registro_xlsx": str(REGISTRO_XLSX),
            "config_path": str(CONFIG_PATH),
            "data_root": str(DATA_ROOT),
            "home": str(HOME),
            "locked_by_env": _automacoes_from_env,
            "default_automacoes_dir": str(DEFAULT_AUTOMAOES_DIR),
        }
    )


@app.post("/api/config")
def api_config_set():
    """Grava pasta das automacoes em config.json (sob Path.home())."""
    if _automacoes_from_env:
        return jsonify(
            {
                "ok": False,
                "error": "SERVERCRON_AUTOMAOES_DIR esta definido no ambiente; remova a env para alterar pelo painel.",
                "automacoes_dir": str(AUTOMAOES_DIR),
            }
        ), 409
    data = request.get_json(silent=True) or {}
    raw = str(data.get("automacoes_dir") or "").strip()
    if not raw:
        return jsonify({"ok": False, "error": "automacoes_dir obrigatorio"}), 400
    try:
        path = _resolve_user_path(raw)
    except OSError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    path.mkdir(parents=True, exist_ok=True)
    apply_automacoes_dir(path, persist=True)
    ensure_automacoes_layout()
    scripts = read_registry(force=True)
    inv = inventory_summary(scripts)
    return jsonify(
        {
            "ok": True,
            "automacoes_dir": str(AUTOMAOES_DIR),
            "registro_xlsx": str(REGISTRO_XLSX),
            "config_path": str(CONFIG_PATH),
            "count": len(scripts),
            "inventory": inv,
        }
    )


@app.get("/api/weather")
def api_weather():
    """Clima atual de Sao Paulo via Open-Meteo (open source, sem chave)."""
    force = request.args.get("force") == "1"
    data = get_weather_sp(force=force)
    status = 200 if data.get("ok") is not False else 502
    return jsonify(data), status


@app.get("/api/scripts")
def api_scripts():
    return jsonify({"ok": True, "scripts": read_registry(force=request.args.get("force") == "1")})


@app.post("/api/reload")
def api_reload():
    """Recarrega planilha Excel e atualiza agendamentos/fila no server."""
    data = reload_spreadsheet_and_schedules()
    return jsonify({"ok": True, **data, "max_concurrent": MAX_CONCURRENT})


@app.get("/api/history")
def api_history():
    limit = int(request.args.get("limit") or 100)
    return jsonify({"ok": True, "history": list_history(limit)})


@app.get("/api/stats")
def api_stats():
    """Tendencia dos ultimos N dias (CSV; exclui catchup)."""
    try:
        days = int(request.args.get("days") or 7)
    except (TypeError, ValueError):
        days = 7
    data = execution_stats(days=days)
    return jsonify({"ok": True, **data})


@app.get("/api/report")
def api_report():
    """Relatorio (aba Relatorio): totais desde o inicio, comparativo semana/mes,
    e drill-down opcional por ano/mes/dia (?year=2026&month=8&day=2)."""
    year = _to_int_or_none(request.args.get("year"))
    month = _to_int_or_none(request.args.get("month"))
    day = _to_int_or_none(request.args.get("day"))
    if month is not None and year is None:
        return jsonify({"ok": False, "error": "year obrigatorio quando month e informado"}), 400
    if day is not None and (year is None or month is None):
        return jsonify({"ok": False, "error": "year e month obrigatorios quando day e informado"}), 400
    return jsonify(report_data(year=year, month=month, day=day))


@app.post("/api/dashboard/rebuild")
def api_dashboard_rebuild():
    """Regenera o Excel a partir do CSV (so leitura do CSV; nao altera o historico)."""
    ok = refresh_excel_dashboard(force=True)
    xlsx = BASE_DIR / "dashboards" / "ServerCRON_Dashboard.xlsx"
    return jsonify(
        {
            "ok": ok,
            "xlsx": str(xlsx) if xlsx.is_file() else None,
            "csv": str(HISTORY_CSV_PATH),
            "csv_modified": False,
            "hint": "Excel regenerado do CSV. Reabra o .xlsx se ja estiver aberto.",
        }
    ), (200 if ok else 500)



@app.post("/api/run")
def api_run():
    if _rate_limited("run", _client_ip(), max_hits=30, window_sec=60):
        return jsonify({"ok": False, "error": "muitas requisicoes, aguarde um pouco"}), 429
    data = request.get_json(silent=True) or {}
    nome = _script_stem(data.get("nome_automacao"))
    if not nome:
        return jsonify({"ok": False, "error": "nome_automacao obrigatorio"}), 400
    scripts = {s["nome_automacao"].lower(): s for s in read_registry(force=True)}
    item = scripts.get(nome.lower())
    if not item:
        return jsonify({"ok": False, "error": "automacao nao encontrada na planilha"}), 404
    if not item.get("path"):
        return jsonify({"ok": False, "error": f"ficheiro {nome}.py nao encontrado"}), 404

    if MANUAL_RUN_COOLDOWN_SEC > 0:
        key = item["nome_automacao"].lower()
        now_ts = time.time()
        with _manual_run_lock:
            last = _manual_run_last.get(key, 0.0)
            if now_ts - last < MANUAL_RUN_COOLDOWN_SEC:
                return jsonify(
                    {
                        "ok": False,
                        "error": "aguarde um instante antes de disparar de novo",
                        "nome_automacao": item["nome_automacao"],
                    }
                ), 429
            _manual_run_last[key] = now_ts

    # Manual tem prioridade maxima na fila (atraso artificial grande)
    ok = enqueue_job(
        item["nome_automacao"],
        item["path"],
        "manual",
        scheduled_for=now_local() - timedelta(days=365),
        area=item.get("area"),
    )
    return jsonify(
        {
            "ok": True,
            "queued": ok,
            "nome_automacao": item["nome_automacao"],
            "running_count": _running_count(),
            "max_concurrent": MAX_CONCURRENT,
            "queue_count": len(list_queue()),
        }
    )


@app.post("/api/kill")
def api_kill():
    if _rate_limited("kill", _client_ip(), max_hits=30, window_sec=60):
        return jsonify({"ok": False, "error": "muitas requisicoes, aguarde um pouco"}), 429
    data = request.get_json(silent=True) or {}
    nome = _script_stem(data.get("nome_automacao"))
    if not nome:
        return jsonify({"ok": False, "error": "nome_automacao obrigatorio"}), 400
    ok = kill_script(nome)
    return jsonify({"ok": ok, "nome_automacao": nome})


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def _open_browser(url: str) -> None:
    time.sleep(1.2)
    try:
        webbrowser.open(url)
    except Exception:
        pass


_dashboard_last_day: str = ""
_dashboard_lock = threading.Lock()


def refresh_excel_dashboard(*, force: bool = False) -> bool:
    """Regenera dashboards/ServerCRON_Dashboard.xlsx (comparativos moveis).

    Janelas (semana atual vs anterior, 30d vs 30d anteriores) usam a data de hoje,
    por isso o ficheiro e recriado 1x por dia automaticamente.
    """
    global _dashboard_last_day
    today = now_local().strftime("%Y-%m-%d")
    with _dashboard_lock:
        if not force and _dashboard_last_day == today:
            return False
        script = BASE_DIR / "dashboards" / "build_excel_dashboard.py"
        if not script.is_file():
            log.warning("[DASHBOARD] script em falta: %s", script)
            return False
        try:
            subprocess.check_call(
                [sys.executable, str(script)],
                cwd=str(script.parent),
                timeout=180,
            )
            _dashboard_last_day = today
            log.info(
                "[DASHBOARD] Excel atualizado (%s) com comparativos moveis",
                today,
            )
            return True
        except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
            log.exception("[DASHBOARD] falha ao regenerar Excel")
            return False


def dashboard_refresh_loop() -> None:
    """Garante rebuild diario do Excel (e na primeira subida do servidor)."""
    # primeira tentativa logo apos o arranque
    _stop_event.wait(8)
    refresh_excel_dashboard(force=True)
    while not _stop_event.is_set():
        # verifica de hora a hora; so regenera quando o dia (SP) muda
        refresh_excel_dashboard(force=False)
        _stop_event.wait(3600)


_LOOPBACK_HOSTS = ("127.0.0.1", "localhost", "::1")


def _local_ips() -> list[str]:
    """Enumera IPs LAN da maquina (wifi, ethernet, etc.) sem depender de libs externas."""
    ips: set[str] = set()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            s.connect(("8.8.8.8", 80))  # nao envia dados, so define a rota/adaptador de saida
            ips.add(s.getsockname()[0])
        finally:
            s.close()
    except OSError:
        pass
    try:
        for ip in socket.gethostbyname_ex(socket.gethostname())[2]:
            if not ip.startswith("127.") and not ip.startswith("169.254."):
                ips.add(ip)
    except OSError:
        pass
    return sorted(ips)


def _guard_network_exposure() -> None:
    """Recusa arrancar exposto na rede sem protecao — evita acesso nao autenticado no LAN."""
    if HOST in _LOOPBACK_HOSTS:
        return
    if AUTH_DISABLED:
        log.warning(
            "[BOOT] SERVERCRON_DISABLE_AUTH ligado: painel exposto em %s SEM NENHUM LOGIN. "
            "Uso temporario para demo -- remova SERVERCRON_DISABLE_AUTH do .env assim que terminar.",
            HOST,
        )
        return
    if _auth_configured():
        log.warning(
            "[BOOT] SERVERCRON_HOST=%s (exposto na rede) — protegido por %s.",
            HOST,
            "SERVERCRON_API_TOKEN" if API_TOKEN else "login por email (aba USERS)",
        )
        return
    log.error(
        "[BOOT] SERVERCRON_HOST=%s expoe o painel na rede sem protecao configurada. "
        "Qualquer pessoa na rede poderia disparar/matar automacoes. "
        "Defina SERVERCRON_API_TOKEN no .env, ou cadastre pelo menos 1 email na aba USERS "
        "de %s, ou volte SERVERCRON_HOST=127.0.0.1.",
        HOST,
        REGISTRO_XLSX,
    )
    raise SystemExit(1)


def main() -> None:
    _guard_network_exposure()
    if read_users() and not _smtp_configured():
        log.error(
            "[BOOT] Aba USERS configurada (login por codigo/email) mas SERVERCRON_SMTP_USER/"
            "SERVERCRON_SMTP_PASSWORD nao estao definidos no .env — NINGUEM vai conseguir logar "
            "(o codigo nao tem como ser enviado). Configure o SMTP ou remova a aba USERS."
        )
    load_automacoes_dir_from_sources()
    ensure_automacoes_layout()
    init_db()
    read_registry(force=True)

    threading.Thread(target=queue_processor_loop, daemon=True, name="queue-processor").start()
    threading.Thread(target=scheduler_loop, daemon=True, name="scheduler").start()
    threading.Thread(target=dashboard_refresh_loop, daemon=True, name="dashboard-xlsx").start()

    url = f"http://{HOST if HOST != '0.0.0.0' else '127.0.0.1'}:{PORT}/"
    log.info("[BOOT] ServerCRON em %s", url)
    if HOST not in _LOOPBACK_HOSTS:
        lan_ips = _local_ips()
        if lan_ips:
            for ip in lan_ips:
                log.info("[BOOT] Acessivel na rede (wifi/ethernet) em: http://%s:%s/", ip, PORT)
        else:
            log.warning("[BOOT] SERVERCRON_HOST=%s mas nao foi possivel detectar o IP da rede automaticamente.", HOST)
    log.info("[BOOT] home=%s", HOME)
    log.info("[BOOT] data_root=%s", DATA_ROOT)
    log.info("[BOOT] Planilha: %s", REGISTRO_XLSX)
    log.info("[BOOT] Automacoes: %s", AUTOMAOES_DIR)
    log.info("[BOOT] SQLite: %s", SQLITE_PATH)
    log.info("[BOOT] Historico CSV: %s", HISTORY_CSV_PATH)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    log.info("[BOOT] max_concurrent=%d catchup_per_script=%d upcoming_per_script=%d", MAX_CONCURRENT, CATCHUP_PER_SCRIPT, UPCOMING_PER_SCRIPT)
    log.info("[BOOT] cron timezone=%s | agora Brasilia=%s", TZ_NAME, now_iso())

    if OPEN_BROWSER:
        threading.Thread(target=_open_browser, args=(url,), daemon=True).start()

    try:
        serve(app, host=HOST, port=PORT, threads=8)
    finally:
        _stop_event.set()


if __name__ == "__main__":
    main()
